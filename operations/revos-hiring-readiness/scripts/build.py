#!/usr/bin/env python3
"""Hiring Readiness Check: write the report (.docx) and the cash workings (.xlsx) from results.json.

Usage:
  python3 scripts/build.py <results.json> --out <folder> [--note note.md] [--role-pack rolepack.json]
        [--md-out report.md] [--format auto|xlsx|csv] [--no-recalc]

The report is the one main file: verdict first, the month-three line, what must be true first, the six
conditions, the cash week by week, the role from the hours it frees, the contractor alternative, the
full-cost checklist and, when the verdict is Ready or the contractor route is chosen, the role pack.
Dates and money follow the founder's locale and currency; nothing here holds a jurisdiction's rules.
The workbook holds the same scenario as live formulas (yellow cells are inputs). Without openpyxl the
workings are written as CSV and the report says so.

The note (at most 120 words) may only quote figures the report already holds; the build refuses it
otherwise. Exit codes: 0 written, 1 written with [NEED] flags, 2 STOP (nothing usable written).
Standard library only (openpyxl optional). No network calls.
"""
import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cash_scenario import (DEFAULTS, STYLE, InputError, cost_who, fmt_plain, long_date, money, pct, set_style,  # noqa: E402
                           symbol)
from docx_writer import write_docx  # noqa: E402

SKILL = "Hiring Readiness Check"
FOOTER = "A planning tool, not financial, tax, legal or employment advice."
# What can go on top of pay, in plain words and with no amounts or rates: a prompt for the founder's conversation with
# their payroll provider or accountant, never an estimate. Jurisdiction-neutral.
ONTOP_ITEMS = ("Pay", "Employer payroll taxes or social charges, where they apply",
               "Insurance an employer carries for staff (cover for injury at work, for example), where it applies",
               "Retirement or pension contributions the employer makes, where they apply",
               "Benefits you offer, such as health cover", "Equipment and software",
               "Cover while they are on paid leave, if you would pay for it", "Training and your own time onboarding them",
               "One-off costs: recruiting, a laptop")
CREDIT = f"Made with Merigrove {SKILL}"  # document properties only, never on a page or sheet


def slug(s):
    s = re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")
    return s[:40] or "role"


def quiet_weeks(days):
    n = int(days // 7)
    words = {1: "week", 2: "two weeks", 3: "three weeks", 4: "four weeks", 5: "five weeks", 6: "six weeks"}
    if n < 1:
        return ""
    if n == 1:
        return ", so the first week brings in nothing"
    return f", so the first {words.get(n, str(n) + ' weeks')} bring in nothing"


def short_date(d):
    return f"{d.strftime('%b')} {d.day}" if STYLE["dates"] in ("mdy", "ymd") else f"{d.day} {d.strftime('%b')}"


def lang_of(locale):
    loc = str(locale or "").replace("_", "-")
    return loc if re.fullmatch(r"[a-z]{2}-[A-Z]{2}", loc) else "en-GB"


def apply_style(inputs):
    i = inputs or {}
    bare = i.get("bare_dollar")
    if bare is None:
        bare = not any(re.search(r"(?<![A-Za-z])[A-Z]{1,3}\$", str(v)) for v in i.values() if isinstance(v, str))
    set_style(i.get("locale"), i.get("country"), i.get("decimal_style"), i.get("currency"), i.get("currency_symbol"), bare)


def bad_quarter_words(b):
    """'20% less money in, every payment 15 days late', or 'paid on time as usual' when the delay is 0."""
    cut, days = b["stress_revenue_cut"], b["stress_extra_days"]
    lead = f"{pct(cut)} less money in" if cut > 0 else "the same money in"
    if b.get("secured_month"):
        lead += f" (contracted revenue of {money(b['secured_month'], b['currency'])} a month not cut)"
    tail = f"every payment {fmt_plain(days)} days late" if days > 0 else "paid on time as usual"
    return f"{lead}, {tail}"


PROS = "your accountant, payroll provider or an employment lawyer where you hire"
STATUS_LINE = ("Whether this person would be an employee or a contractor is not something this check decides. "
               f"Ask {PROS} before you agree terms.")
RULES_LINE = ("Before you post it, check the employment and anti-discrimination rules where you hire; "
              "an employment lawyer or adviser there can confirm the wording.")


def months_words(x):
    words = {1: "one month", 2: "two months", 3: "three months", 6: "six months"}
    return words.get(x, f"{fmt_plain(x, 1)} months") if float(x).is_integer() else f"{fmt_plain(x, 1)} months"


def cell(x):
    return str(x).replace("|", "\\|").replace("\n", " ")


# ------------------------------------------------------------------ the report
def report_md(res, note=None, rp=None, csv_fallback=False, lead=None, lead_above=False):
    cur = res["currency"]
    m = lambda v: money(v, cur)  # noqa: E731
    b = res["inputs"]
    role = res["role"]
    today = date.fromisoformat(res["date"])
    start = date.fromisoformat(role["start"])
    sc = res.get("scenario")
    L = []
    L.append(f"# Hiring readiness: {role['title']}")
    named = res["business"] if res["business"] != "Your business" else ""
    L.append(f"_{named + '. ' if named else ''}{long_date(today)}. A planning tool, not advice: see the last section._")
    lead_lines = [f"**{x.strip()}**" for x in (lead or "").splitlines() if x.strip()]
    if lead_above:
        L += lead_lines
    L.append(f"VERDICT: {res['verdict']}.")
    L.append(res["verdict_reason"])
    if not lead_above:
        L += lead_lines
    if sc:
        line = (f"**Month three of the hire:** in week {sc['end_week']} (the week of {long_date(date.fromisoformat(sc['end_date']))}), the hire's 13th week, the bank holds "
                f"{m(sc['base_end'])} if the next three months look like the last three, and {m(sc['stress_end'])} in a bad quarter "
                f"({bad_quarter_words(b)}). "
                f"Your floor is {m(sc['floor'])}: {months_words(b['cover_target_months'])} of costs with the hire in them.")
        if sc["stress_breach"]:
            line += f" In the bad quarter the bank drops below the floor in week {sc['stress_breach']}."
        else:
            line += " In the bad quarter the bank stays above the floor in every week."
        if role.get("cost_basis") == "salary_only":
            line += " **These figures carry pay alone, so they understate the cost of the hire.**"
        L.append(line)
    else:
        L.append("**Month three:** [NEED: cash today, monthly revenue, monthly costs and the hire's full monthly cost] before the bank balance can be worked out.")
    ce = res.get("ceiling")
    if ce:
        if ce["usual"] <= 0:
            L.append("**What the business can carry:** nothing yet. Even without the hire, the usual quarter drops below the floor.")
        else:
            line = (f"**What the business can carry:** the hire can cost up to {m(ce['usual'])} a month, all in, before the usual quarter "
                    "drops below the floor by the hire's 13th week; "
                    + (f"up to {m(ce['bad'])} in a bad quarter." if ce["bad"] > 0 else
                       "in a bad quarter there is no room, because the floor breaks even without the hire."))
            if ce.get("cost_used") is not None and ce["basis"] == "salary_only":
                line += (f" Pay alone is {m(ce['cost_used'])}, so whatever goes on top of pay has {m(ce['room_usual'])} a month of room "
                         "in a usual quarter. When the fully loaded figure arrives, compare it with these.")
            elif ce.get("cost_used") is not None:
                over = [k for k in ("usual", "bad") if ce[k] > 0 and ce["cost_used"] > ce[k] + 1e-9]
                line += (f" The {m(ce['cost_used'])} used is inside both." if not over and ce["bad"] > 0 else
                         f" The {m(ce['cost_used'])} used is over the {' and '.join(k for k in over)} quarter figure." if over else "")
            L.append(line)
    pw_ = res.get("pays_its_way")
    if pw_:
        line = (f"**What it has to bring in:** at {pct(pw_['margin'])} gross margin, the hire covers its own cost when it brings in or "
                f"protects {m(pw_['revenue_month'])} a month of revenue"
                + (" (on pay alone, so the real figure is higher)" if pw_["basis"] == "salary_only" else ""))
        if pw_.get("clients"):
            line += f": {pw_['clients']} new clients at {m(pw_['typical_client_month'])} a month"
        L.append(line + ".")
    if note:
        L.append("> " + " ".join(note.split()))

    if res["verdict"] != "Ready":
        L.append("## What would have to change" if res["verdict"] == "No" else "## What must be true first")
        k = 0
        for c in res["conditions"]:
            if c["status"] == "Pass":
                continue
            for x in c["must_be_true"] or [f"[NEED: {', '.join(c['needs']) or 'the figures for this condition'}]"]:
                k += 1
                L.append(f"{k}. {x}. ({c['name']})")
        if res["role"].get("cost_basis") != "loaded" and res["role"].get("monthly_cost") is not None:
            k += 1
            L.append(f"{k}. The hire's fully loaded monthly cost from {cost_who(b)} "
                     f"(the {m(res['role']['monthly_cost'])} used is {'pay alone' if res['role'].get('cost_basis') == 'salary_only' else 'not confirmed as fully loaded'}).")
        L.append("These are conditions, not a date. When they hold, run the check again with that month's figures.")

    L.append("## The six conditions")
    L.append("| # | Condition | Result | Evidence |")
    L.append("|---|---|---|---|")
    for c in res["conditions"]:
        L.append(f"| {c['n']} | {cell(c['name'])} | {c['status']} | {cell(c['evidence'])} |")
    passed = sum(1 for c in res["conditions"] if c["status"] == "Pass")
    words = ["None", "One", "Two", "Three", "Four", "Five", "All six"]
    L.append(f"{words[passed]} of six pass. Ready needs all six; a condition with a missing figure cannot pass."
             if passed < 6 else ("All six pass." if role.get("cost_basis") == "loaded" else
                                 "All six pass, but on a cost that is not the fully loaded figure, so the verdict cannot be Ready."))

    if sc:
        L.append("## Your cash, week by week")
        surplus = b["monthly_revenue"] - b["costs_month_total"]
        L.append(f"Before any hire, {m(b['monthly_revenue'])} a month comes in and {m(b['costs_month_total'])} goes out: "
                 + (f"{m(surplus)} a month left over." if surplus >= 0 else f"{m(-surplus)} a month short."))
        if b["stress_extra_days"] > 0:
            how = (f"In the bad quarter every payment arrives {fmt_plain(b['stress_extra_days'])} days later and brings in "
                   f"{pct(b['stress_revenue_cut'])} less" + quiet_weeks(b['stress_extra_days']))
        else:
            how = f"In the bad quarter {pct(b['stress_revenue_cut'])} less comes in, paid on time as usual"
        if b.get("variable_costs_share"):
            how += f"; costs fall by {pct(b['variable_costs_share'])} of the revenue lost, your figure for the costs that move with revenue"
        if b.get("secured_month"):
            how += (f"; your contracted revenue of {m(b['secured_month'])} a month is not cut"
                    + (", but arrives just as late" if b["stress_extra_days"] > 0 else ""))
        L.append(f"The hire costs {m(role['monthly_cost'])} a month from {long_date(start)} (week {role['start_week']})"
                 + (f", plus {m(role['one_off'])} of one-off costs that week" if role["one_off"] else "") + ". " + how + ".")
        if sc.get("nohire_shortfall", 0) > 0.5:
            L.append(f"Without the hire, the same bad quarter already falls {m(sc['nohire_shortfall'])} short of the floor "
                     f"(week {sc['nohire_shortfall_week']}). The hire adds {m(sc['hire_adds'])} to that.")
        for lp in b.get("low_periods") or []:
            L.append(f"Known low period loaded in both cases: {lp['what']}, {long_date(date.fromisoformat(lp['from']))} to "
                     f"{long_date(date.fromisoformat(lp['to']))}, "
                     + ("nothing coming in" if lp["revenue_cut"] >= 0.999 else f"{pct(lp['revenue_cut'])} less coming in")
                     + (f" and {pct(lp['costs_cut'])} less going out" if lp["costs_cut"] else "") + ".")
        if sc.get("window_weeks", 0) > sc.get("end_week", 0):
            L.append(f"The table runs on to week {sc['window_weeks']} to take in the low period; month three of the hire is week {sc['end_week']}.")
        L.append("| Week | Week of | Usual quarter | Bad quarter | Floor |")
        L.append("|---|---|---|---|---|")
        for bw, sw in zip(sc["base"], sc["stress"]):
            w = bw["week"]
            fl = sc["floor"] if w >= role["start_week"] else sc["floor_before_start"]
            mark = " (below)" if sw["balance"] < fl - 1e-9 else ""
            L.append(f"| {w} | {short_date(today + timedelta(days=7 * (w - 1)))} | {m(bw['balance'])} | {m(sw['balance'])}{mark} | {m(fl)} |")
        kp = [p for p in b["known_payments"] if p["in_window"]]
        if kp:
            L.append("Known bills loaded: " + "; ".join(f"{p['what']} {m(p['amount'])} in week {p['week']}" for p in kp) + ".")
        seas = b.get("season") or []
        if seas:
            L.append("**Seasons the flat average may hide.** " + " ".join(s["text"] + "." for s in seas)
                     + " If a quiet spell repeats, give its dates and roughly how much less comes in, and the check loads it.")

    L.append("## The role, built from the hours it frees")
    if b["tasks"]:
        L.append("| Task it takes off you | Hours a week | Written down |")
        L.append("|---|---|---|")
        for t in b["tasks"]:
            L.append(f"| {cell(t['task'])} | {fmt_plain(t['hours_week'] or 0, 1)} | {'not said' if t['written_down'] == 'unknown' else t['written_down']} |")
        th = sum(t["hours_week"] or 0 for t in b["tasks"])
        L.append(f"Named tasks come to {fmt_plain(th, 1)} hours a week against a role of "
                 f"{fmt_plain(role['hours_week'], 1) if role['hours_week'] else '[NEED]'} hours.")
    else:
        L.append("[NEED: the tasks this person would take off you, with rough hours a week. A week of your calendar is enough.]")
    L.append(f"**Your freed hours go to:** {b['freed_hours_use'] or '[NEED: what you would do with them]'}.")
    L.append(res["order_note"])

    L.append("## The contractor alternative")
    ct = res["contractor"]
    if ct["status"] == "run":
        for r in ct["runs"]:
            verdict = ("passes the cash test today" if r["passes"]
                       else f"fails the cash test: the bad quarter drops below the floor in week {r['stress_breach']}")
            L.append(f"- {fmt_plain(r['hours_week'], 1)} hours a week at {m(ct['rate'])} an hour is {m(r['monthly'])} a month and {verdict}. "
                     f"Bad quarter, the hire's 13th week: {m(r['stress_end'])}.")
        L.append(f"A contractor is a cost you can stop. Under {fmt_plain(ct['line'])} hours a week, or where the work can be defined "
                 "piece by piece, start there; an employee suits continuous work that needs your judgement passed on over months.")
    else:
        L.append(f"Not run: [NEED: {ct['need']}]. A contractor is a cost you can stop; under {fmt_plain(ct['line'])} hours a week it is the usual first step.")
    L.append(STATUS_LINE)

    L.append("## The full cost of the hire")
    mc = role.get("monthly_cost")
    who = role.get("cost_from") or "you"
    if mc is None:
        L.append("The scenario has no cost for the hire yet: [NEED: the fully loaded monthly cost, pay plus whatever your accountant "
                 "or payroll provider says goes on top of it where you hire]. This check never works it out.")
    elif role.get("cost_basis") == "loaded":
        L.append(f"The scenario uses {m(mc)} a month, the fully loaded cost from {who}. This check never works it out. "
                 "Check it covers each of these:")
    elif role.get("cost_basis") == "salary_only":
        L.append(f"**The scenario uses {m(mc)} a month, which is pay alone. It understates the cost of the hire**, because whatever goes on "
                 f"top of pay where you hire is missing, and the verdict cannot be Ready on it. Ask {cost_who(b)} "
                 "for the fully loaded monthly cost and run the check again as soon as you have it. Check their figure covers each of these:")
    else:
        L.append(f"The scenario uses {m(mc)} a month, which is not confirmed as the fully loaded cost, so the verdict cannot be Ready on it. "
                 f"Check with {cost_who(b)} that it covers each of these:")
    for item in ONTOP_ITEMS:
        L.append(f"- {item}")
    L.append(f"Which of these apply, and at what amounts, depends on where and how you hire; this check never works them out. "
             f"Ask {cost_who(b)} for one monthly total that covers them.")
    if role.get("one_off"):
        L.append(f"One-off costs of {m(role['one_off'])} (recruitment, equipment) are loaded in the start week.")
    if role.get("hours_conflict"):
        L.append(f"**Check the hours before you rely on this cost.** You said {fmt_plain(role['hours_week'], 1)} hours a week and "
                 f"'{role['days']}', which is {fmt_plain(role['hours_conflict'], 1)} hours at {fmt_plain(STYLE.get('day_hours', 8), 1)} "
                 f"a day. The {m(mc)} a month is for one of them; ask {role.get('cost_from') or cost_who(b)} which, and run the check "
                 "again if it changes.")

    if rp:
        L += role_pack_md(rp, res)

    ch = res.get("changes")
    if ch and ch["lines"]:
        L.append(f"## Since {long_date(date.fromisoformat(ch['since'])) if ch.get('since') else 'last time'}")
        for x in ch["lines"]:
            L.append(f"- {x}")

    L.append("## Assumptions and sources")
    for a in res["assumptions"]:
        L.append(f"- Assumed: {a}.")
    for n in res["notes"]:
        L.append(f"- {n}.")
    srcs = b.get("sources") or {}
    labels = {"role_monthly": "The hire's monthly cost", "known_payments": "Known bills", "contractor_rate": "Contractor rate",
              "contractor_rate_hour": "Contractor rate", "contractor_line_hours": "Contractor line", "revenue_history": "Revenue by month",
              "margin_detail": "Gross margin", "gross_margin_months": "Gross margin", "salary_year": "Salary",
              "cash_today": "Cash today", "monthly_revenue": "Monthly revenue", "monthly_costs": "Monthly costs"}
    for k, v in srcs.items():
        L.append(f"- {labels.get(k, k.replace('_', ' ').capitalize())}: {v}.")
    for n in res["needs"]:
        L.append(f"- [NEED: {n}]")
    if csv_fallback:
        L.append("- The cash workings are saved as CSV files because the spreadsheet library was not available; the figures are the same.")
    L.append("## Before you act")
    L.append("This is a planning tool, not financial, tax, legal or employment advice. Every figure comes from the numbers you gave and the stated defaults. "
             f"What goes on top of pay, payroll, contracts and whether this person would be an employee or a contractor: {PROS}.")
    return "\n".join(L) + "\n"


def role_pack_md(rp, res):
    L = []
    emp = str(rp.get("employment", "")).lower()
    title = res["role"]["title"]
    verdict = res.get("verdict")
    hold = None
    if verdict and verdict not in ("Ready", "not checked"):
        ct = res.get("contractor") or {}
        con_ok = ct.get("status") == "run" and any(r["passes"] for r in ct.get("runs", []))
        if emp == "contractor":
            if not con_ok:
                hold = ("Written at your request. This contractor's cost has not passed the cash test yet; "
                        "agree a fixed first piece of work before any rolling arrangement.")
        else:
            hold = ("Written at your request. The conditions above are not met yet; hold the ad until they are"
                    + (", or use the contractor route, which passes the cash test today." if con_ok else "."))
    if emp == "contractor":
        L.append("## The contractor brief")
        if hold:
            L.append(f"**{hold}**")
        L.append(f"**Mission:** {rp.get('mission', '')}")
        L.append("Managed by output, not hours. Each piece is done when:")
        L.append("| Deliverable | Done when |")
        L.append("|---|---|")
        for d in rp.get("contractor_deliverables") or []:
            L.append(f"| {cell(d.get('deliverable', ''))} | {cell(d.get('done_when', ''))} |")
    else:
        L.append(f"## The role pack: {title}")
        if hold:
            L.append(f"**{hold}**")
        L.append("### Scorecard")
        L.append(f"**Mission:** {rp.get('mission', '')}")
    if emp == "contractor":
        L.append("### Outcomes")
    L.append("| # | Outcome | Measure | By | From your |")
    L.append("|---|---|---|---|---|")
    for i, o in enumerate(rp.get("outcomes") or [], 1):
        frm = f"{o.get('from_task', '')}" + (f" ({fmt_plain(float(o['hours_week']), 1)} h a week)" if o.get("hours_week") else "")
        L.append(f"| {i} | {cell(o.get('outcome', ''))} | {cell(o.get('measure', ''))} | {cell(o.get('by', ''))} | {cell(frm)} |")
    comp = rp.get("competencies") or []
    if comp:
        L.append("**What they need to be good at:** " + "; ".join(f"{c.get('competency')} (outcome {c.get('for_outcome')})" for c in comp) + ".")
    ad = rp.get("ad") or {}
    L.append("### What to send the agency or freelancer" if emp == "contractor" else "### Job ad")
    if ad.get("headline"):
        L.append(f"**{ad['headline']}**")
    for k in ("about", "what_you_will_do"):
        if ad.get(k):
            L.append(ad[k])
    if ad.get("what_you_will_own"):
        L.append("**What you will own**")
        for x in ad["what_you_will_own"]:
            L.append(f"- {x}")
    if ad.get("good_at_90_days"):
        L.append(f"**What good looks like at day 90:** {ad['good_at_90_days']}")
    L.append(f"**Pay:** {ad.get('pay') or '[Not shown: ' + (ad.get('pay_note') or 'your choice') + ']'}. **Hours:** {ad.get('hours', '')}. **Location:** {ad.get('location', '')}.")
    if ad.get("how_to_apply"):
        L.append(f"**How to apply:** {ad['how_to_apply']}")
    L.append(f"_{RULES_LINE}_")
    L.append("### Interview guide")
    L.append("| # | Question | A strong answer includes |")
    L.append("|---|---|---|")
    for q in rp.get("interview") or []:
        L.append(f"| {q.get('outcome')} | {cell(q.get('question', ''))} | {cell(q.get('strong_answer', ''))} |")
    ws = rp.get("work_sample") or {}
    if ws.get("task"):
        L.append(f"**Work sample ({ws.get('time', '')}, {ws.get('paid', 'founder to decide')}):** {ws['task']}")
    if emp != "contractor":
        L.append("### The first 90 days")
        L.append("| Weeks | Focus | They decide |")
        L.append("|---|---|---|")
        for s in rp.get("ninety_days") or []:
            L.append(f"| {cell(s.get('weeks', ''))} | {cell(s.get('focus', ''))} | {cell(s.get('decision_rights', ''))} |")
        if rp.get("handover_test"):
            L.append(rp["handover_test"])
        if rp.get("one_to_one"):
            L.append("### Fortnightly one-to-one (30 minutes, their agenda)")
            for x in rp["one_to_one"]:
                L.append(f"- {x}")
    if emp == "contractor":
        L.append(f"Before they start: a written agreement with the agency or freelancer. {STATUS_LINE}")
    else:
        L.append(f"Before an offer: the contract, payroll set-up and any checks you make on the person are for {PROS} to confirm.")
    return L


# ------------------------------------------------------------------ note gate
TOKEN_RE = re.compile(r"(?<![\w.,$])(-?)((?:[A-Z]{1,3})?[$\u00a3\u20ac\u20b9\u00a5\u20b1])?(\d[\d.,]*\d|\d)(\s?%| per cent|\s?[\u20ac\u00a3$])?")
PROMISES = re.compile(r"\b(you will make|you'll make|guarantee[sd]?|guaranteed|risk[- ]free|you can afford|ready by)\b", re.I)
RATIOS = re.compile(r"\b(twice|(?:one|two|three|four|five|six|seven|eight|nine|ten|\d+(?:\.\d+)?) times( over)?|half (?:of|as)|double|triple|a third of|a quarter of)\b", re.I)
SMALL_UNITS = r"(st|nd|rd|th)?\s*(hours?|h\b|weeks?|months?|days?|clients?|calls?|people|person|tasks?|conditions?|of six|of the six|proposals?)"


def tokens(text, style=None):
    """Numbers in text, read in the founder's style (dot: 1,234.5; comma: 1.234,5)."""
    out = []
    for mm in TOKEN_RE.finditer(text):
        n = mm.group(3)
        n = n.replace(".", "").replace(",", ".") if style == "comma" else n.replace(",", "")
        try:
            v = round(float(mm.group(1) + n), 2)
        except ValueError:
            continue
        suf = mm.group(4) or ""
        kind = "money" if (mm.group(2) or suf.strip() in ("\u20ac", "\u00a3", "$")) else ("pct" if suf.strip() in ("%", "per cent") else "plain")
        out.append((v, kind, mm.end()))
    return out


def money_values(res):
    """Money figures the note may quote with a currency sign: the founder's money inputs, the scenario's figures
    and the contractor runs."""
    b, out = res["inputs"], []
    for k in ("cash_today", "monthly_revenue", "monthly_costs", "costs_month_total", "founder_pay_month",
              "recurring_revenue_month", "contracted_next_6_months", "pipeline_weighted", "contractor_rate_hour"):
        out.append(b.get(k))
    r = res["role"]
    out += [r.get("monthly_cost"), r.get("salary_year"), r.get("one_off")]
    out += [p["amount"] for p in b.get("known_payments") or []] + [d["value"] for d in b.get("pipeline") or []]
    sc = res.get("scenario") or {}
    out += [sc.get(k) for k in ("floor", "floor_before_start", "shortfall", "base_end", "stress_end", "stress_low")]
    out += [x["balance"] for x in sc.get("base", []) + sc.get("stress", [])]
    for run in (res.get("contractor") or {}).get("runs", []):
        out += [run["monthly"], run["stress_end"], run["base_end"], run["floor"]]
    vals = []
    for v in out:
        if isinstance(v, (int, float)):
            vals += [round(v, 2), round(v)]
    return vals


def check_note(note, allowed_text, style=None):
    """allowed_text: text the figures may come from (read in the founder's style), or a list of (value, kind) pairs."""
    allowed = {}
    pairs = tokens(allowed_text, style) if isinstance(allowed_text, str) else allowed_text
    for item in pairs:
        v, k = item[0], item[1]
        allowed.setdefault(k, set()).add(v)
        allowed.setdefault("any", set()).add(v)
    probs = []
    for v, k, end in tokens(note, style):
        if k == "plain" and 0 <= v <= 13 and float(v).is_integer() and re.match(SMALL_UNITS, note[end:], re.I):
            continue
        if k == "plain" and 16 <= v <= 99 and float(v).is_integer() and re.match(r"\s*(and over|or over|years? old|year-olds?)", note[end:], re.I):
            continue
        ok = v in allowed.get(k, set()) if k in ("money", "pct") else v in allowed.get("any", set())
        if not ok:
            shown = f"{v:,.2f}".rstrip("0").rstrip(".") if not float(v).is_integer() else fmt_plain(v)
            probs.append(f"figure {shown}{'%' if k == 'pct' else ''} is not in the results" + (" as a money figure" if k == "money" else ""))
    for mm in RATIOS.finditer(note):
        probs.append(f"ratio worked out in prose: '{mm.group(0)}'; quote the two figures instead")
    for mm in PROMISES.finditer(note):
        probs.append(f"promise: '{mm.group(0)}'")
    words = len(re.findall(r"\b\w[\w'-]*\b", note))
    if words > 120:
        probs.append(f"note is {words} words, limit 120")
    if re.search("[\u2014\u2013]", note):
        probs.append("note contains an em or en dash")
    return probs


# ------------------------------------------------------------------ workbook
def recalc(path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return False
    tmp = Path(tempfile.mkdtemp())
    try:
        subprocess.run([exe, f"-env:UserInstallation=file://{tmp / 'profile'}", "--headless", "--calc",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "out"), str(path)],
                       check=True, capture_output=True, timeout=120)
        made = tmp / "out" / Path(path).name
        if made.exists() and made.stat().st_size > 0:
            shutil.copyfile(made, path)
            return True
    except Exception:  # noqa: BLE001
        return False
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return False


IN = {  # key: (row, label, kind)
    "date": (3, "Date of this run (week 1 starts here)", "date"),
    "cash_today": (4, "Cash in the bank today", "money"),
    "monthly_revenue": (5, "Average monthly revenue, last three months", "money"),
    "costs_month_total": (6, "Average monthly costs, your own pay included", "money"),
    "repeat_month": (7, "Repeat revenue a month (regulars, memberships, plans)", "money"),
    "stress_revenue_cut": (8, "Bad quarter: money in lower by", "pct"),
    "stress_extra_days": (9, "Bad quarter: every payment this many days later", "num"),
    "cover_target_months": (10, "Cash floor, in months of costs (hire included)", "num"),
    "role_monthly": (11, "The hire's monthly cost (your figure)", "money"),
    "one_off": (12, "One-off costs in the start week", "money"),
    "start_week": (13, "Start week (week 1 is this week)", "num"),
    "contractor_rate": (14, "Contractor rate an hour (a quote)", "money"),
    "role_hours": (15, "Role hours a week", "num"),
    "contractor_line": (16, "Contractor line, hours a week", "num"),
    "margin_floor": (17, "Gross margin floor", "pct"),
    "gm1": (18, "Gross margin, month before last", "pct"),
    "gm2": (19, "Gross margin, last month", "pct"),
    "recurring_month": (20, "Recurring revenue a month", "money"),
    "contracted_6": (21, "Contracted revenue, next six months (leave empty to use recurring x 6)", "money"),
    "pipeline_weighted": (22, "Weighted pipeline", "money"),
    "cap_now": (23, "How full you are now (client hours or clients)", "num"),
    "cap_max": (24, "The most you can take on", "num"),
    "near_ceiling": (25, "Near full means at least", "pct"),
    "price_rise": (26, "Price rise tested for new clients (yes or no)", "text"),
    "employment": (27, "Employee or contractor", "text"),
    "cost_confirmed": (28, "Hire cost is the fully loaded figure from your accountant or payroll provider (yes or no)", "text"),
    "secured_month": (29, "Contracted revenue a month not cut in a bad quarter (still paid late)", "money"),
    "variable_costs_share": (42, "Bad quarter: costs fall by this share of the revenue lost (costs that move with revenue)", "pct"),
}
KP_ROW = 31
LP_ROW = 44   # known low periods: header row, then three rows
HIST_ROW = 50  # 24 months, newest in the last row
HIST_N = 24
SRC_KEYS = {"costs_month_total": ("monthly_costs", "costs_month_total"), "role_monthly": ("role_monthly", "monthly_cost", "role", "salary_year"),
            "contractor_rate": ("contractor_rate_hour", "contractor_rate"), "contractor_line": ("contractor_line_hours",),
            "gm1": ("gross_margin_months", "margin_detail", "gross_margin_rough"), "gm2": ("gross_margin_months", "margin_detail", "gross_margin_rough"),
            "recurring_month": ("recurring_revenue_month",), "contracted_6": ("contracted_next_6_months",),
            "repeat_month": ("repeat_revenue_month", "repeat_share"), "secured_month": ("contracted_next_6_months",),
            "stress_revenue_cut": ("stress_revenue_cut",), "stress_extra_days": ("stress_extra_days",),
            "variable_costs_share": ("variable_costs_share",),
            "cover_target_months": ("cover_target_months",), "price_rise": ("price_rise_tested",), "cap_now": ("capacity",),
            "cap_max": ("capacity",), "start_week": ("start", "role_start"), "role_hours": ("hours_week", "role_hours")}


def write_xlsx(path, res):
    import openpyxl  # optional: build.py falls back to CSV without it
    from openpyxl.styles import Alignment, Font, PatternFill
    b = res["inputs"]
    role = res["role"]
    cur = res["currency"]
    sym = symbol(cur)
    after = sym and cur in ("EUR", "CHF", "SEK", "NOK", "DKK", "PLN", "CZK", "HUF") and STYLE.get("decimal") == "comma"
    mfmt = (f'#,##0 "{sym}";-#,##0 "{sym}"' if after else f'"{sym}"#,##0;-"{sym}"#,##0') if sym else "#,##0"
    dfmt = {"mdy": "mmm d, yyyy", "ymd": "yyyy-mm-dd"}.get(STYLE["dates"], "dd mmm yyyy")
    yellow = PatternFill("solid", fgColor="FFF4C2")
    bold = Font(bold=True)
    pay_alone = role.get("cost_basis") == "salary_only"
    wb = openpyxl.Workbook()
    rd = wb.active
    rd.title = "Read me"
    rd["A1"] = f"Hiring readiness: {role['title']}"
    rd["A1"].font = Font(bold=True, size=14)
    lines = [(f"{res['business']}, run" if res["business"] != "Your business" else "Run") + f" on {long_date(date.fromisoformat(res['date']))}.",
             "Yellow cells on Inputs are your figures. Change one and the cash weeks, the floor and the conditions recalculate.",
             "Cash by week: from today to the hire's 13th week (and on through any low period you named), in a usual quarter and a bad one "
             "(less money in except contracted revenue, every payment later), the hire loaded from the start week. A start later than the one you gave needs a re-run.",
             "Conditions: the six tests. Condition 5 (a real role) is a judgement, carried from the report.",
             "Contractor: the same bad quarter with a contractor in place of the hire."]
    if pay_alone:
        lines.append("The hire's cost on Inputs is PAY ALONE. It understates the hire: whatever goes on top of pay where you hire is "
                     f"missing. Ask {cost_who(b)} for the fully loaded monthly figure and put it in its place.")
    lines += [f"Planning, not advice. What goes on top of pay, payroll, contracts and employee or contractor status: {PROS}."]
    for i, txt in enumerate(lines, 3):
        rd.cell(row=i, column=1, value=txt)
        if pay_alone and txt.startswith("The hire's cost on Inputs is PAY ALONE"):
            rd.cell(row=i, column=1).font = bold
    rd.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Inputs")
    ws["A1"] = "Your figures"
    ws["A1"].font = Font(bold=True, size=13)
    ws["C2"] = "Where it came from"
    ws["C2"].font = bold
    gm = b["gross_margin_months"]
    rough = b.get("gross_margin_rough")
    if len(gm) < 2 and rough:
        gm1, gm2 = rough["low"], rough["high"]
    else:
        gm1, gm2 = (gm[0] if len(gm) > 1 else None), (gm[-1] if gm else None)
    vals = {"date": date.fromisoformat(res["date"]), "cash_today": b["cash_today"], "monthly_revenue": b["monthly_revenue"],
            "costs_month_total": b["costs_month_total"], "repeat_month": b.get("repeat_revenue_month"),
            "stress_revenue_cut": b["stress_revenue_cut"],
            "stress_extra_days": b["stress_extra_days"], "cover_target_months": b["cover_target_months"],
            "role_monthly": role["monthly_cost"], "one_off": role["one_off"], "start_week": role["start_week"],
            "contractor_rate": b["contractor_rate_hour"], "role_hours": role["hours_week"],
            "contractor_line": b["contractor_line_hours"], "margin_floor": b["margin_floor"],
            "gm1": gm1, "gm2": gm2,
            "recurring_month": b["recurring_revenue_month"], "contracted_6": b["contracted_next_6_months"],
            "pipeline_weighted": b["pipeline_weighted"], "cap_now": b["capacity"]["now"], "cap_max": b["capacity"]["max"],
            "near_ceiling": DEFAULTS["near_ceiling"][0],
            "price_rise": {True: "yes", False: "no"}.get(b["price_rise_tested"]["value"], ""),
            "employment": role["employment"],
            "cost_confirmed": "yes" if role.get("cost_basis") == "loaded" else "no",
            "secured_month": b.get("secured_month") or 0.0,
            "variable_costs_share": b.get("variable_costs_share") or 0.0}
    labels = {k: v[1] for k, v in IN.items()}
    if pay_alone:
        labels["role_monthly"] = "The hire's monthly cost (PAY ALONE: understates the hire)"
    if len(gm) < 2 and rough:
        labels["gm1"] = f"Gross margin, low end of your rough range ({rough['said']})"
        labels["gm2"] = "Gross margin, high end of your rough range"
    srcs = b.get("sources") or {}
    dflt = set(b.get("defaults_used") or [])
    for k, (r, _lab, kind) in IN.items():
        ws.cell(row=r, column=1, value=labels[k])
        c = ws.cell(row=r, column=2, value=vals.get(k))
        c.fill = yellow
        if kind == "money":
            c.number_format = mfmt
        elif kind == "pct":
            c.number_format = "0.0%"
        elif kind == "date":
            c.number_format = dfmt
        src = next((srcs[x] for x in (k,) + SRC_KEYS.get(k, ()) if srcs.get(x)), "")
        if k == "role_monthly":
            src = src or (f"from {role.get('cost_from')}" if role.get("cost_from") else "")
            if pay_alone:
                src = (src + "; " if src else "") + "pay alone"
        if not src and (k in dflt or k in ("near_ceiling",)):
            src = "default"
        ws.cell(row=r, column=3, value=src)
    if pay_alone:
        ws.cell(row=IN["role_monthly"][0], column=1).font = bold
    ws.cell(row=KP_ROW - 1, column=1, value="Known bills (tax payments, renewals, annual bills)").font = bold
    for j, h in enumerate(("What", "Amount", "Date", "Week"), 1):
        ws.cell(row=KP_ROW, column=j, value=h).font = bold
    kps = b["known_payments"]
    for i in range(10):
        r = KP_ROW + 1 + i
        p = kps[i] if i < len(kps) else None
        for j in (1, 2, 3):
            ws.cell(row=r, column=j).fill = yellow
        if p:
            ws.cell(row=r, column=1, value=p["what"])
            ws.cell(row=r, column=2, value=p["amount"]).number_format = mfmt
            ws.cell(row=r, column=3, value=date.fromisoformat(p["date"])).number_format = dfmt
        ws.cell(row=r, column=4, value=f'=IF(C{r}="","",MAX(1,INT((C{r}-$B$3)/7)+1))')
    ws.cell(row=LP_ROW - 1, column=1, value="Known low periods: a closure or quiet season, with how much less comes in and goes out").font = bold
    for j, h in enumerate(("What", "From", "To (last day)", "Money in lower by", "Costs lower by", "Starts (weeks from today)", "Ends (weeks from today)"), 1):
        ws.cell(row=LP_ROW, column=j, value=h).font = bold
    lps = b.get("low_periods") or []
    for i in range(3):
        r = LP_ROW + 1 + i
        p = lps[i] if i < len(lps) else None
        for j in (1, 2, 3, 4, 5):
            ws.cell(row=r, column=j).fill = yellow
        if p:
            ws.cell(row=r, column=1, value=p["what"])
            ws.cell(row=r, column=2, value=date.fromisoformat(p["from"])).number_format = dfmt
            ws.cell(row=r, column=3, value=date.fromisoformat(p["to"])).number_format = dfmt
            ws.cell(row=r, column=4, value=p["revenue_cut"]).number_format = "0%"
            ws.cell(row=r, column=5, value=p["costs_cut"]).number_format = "0%"
        ws.cell(row=r, column=6, value=f'=IF(B{r}="",0,MAX(0,(B{r}-$B$3)/7))')
        ws.cell(row=r, column=7, value=f'=IF(C{r}="",0,MAX(0,(C{r}-$B$3+1)/7))')
    ws.cell(row=HIST_ROW - 1, column=1, value="Revenue by month, oldest first, newest in the last row (up to 24 months)").font = bold
    hist = b["revenue_history"][-HIST_N:]
    labs = (b.get("revenue_months") or [])[-HIST_N:]
    pad = HIST_N - len(hist)
    for i in range(HIST_N):
        r = HIST_ROW + i
        j = i - pad
        ws.cell(row=r, column=1, value=(labs[j] if labs and j >= 0 else f"Month {i + 1}"))
        c = ws.cell(row=r, column=2, value=hist[j] if j >= 0 else None)
        c.fill = yellow
        c.number_format = mfmt
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 40

    I = lambda k: f"Inputs!$B${IN[k][0]}"  # noqa: E731
    cs = wb.create_sheet("Cash by week")
    heads = ["Week", "Week of", "In, usual", "In, bad quarter", "Running costs", "The hire", "Known bills",
             "Bank, usual", "Bank, bad quarter", "Floor", "Bad quarter below floor?",
             "Contractor cost", "Bank, bad quarter, contractor", "Floor with contractor",
             "Contractor at the line", "Bank, bad quarter, contractor at the line", "Floor, contractor at the line"]
    for j, h in enumerate(heads, 1):
        c = cs.cell(row=4, column=j, value=h)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
    cs["A1"] = "Cash, week by week, with the hire loaded"
    cs["A1"].font = Font(bold=True, size=13)
    cs["A2"] = "Opening balance"
    cs["B2"] = f"={I('cash_today')}"
    cs["B2"].number_format = mfmt
    W = f"({I('monthly_revenue')}*12/52)"
    S = f"(MIN(N({I('secured_month')}),{I('monthly_revenue')})*12/52)"
    D = f"({I('stress_extra_days')}/7)"
    CUT = I("stress_revenue_cut")
    REL = f"(N({I('variable_costs_share')})*({W}-{S})*{CUT})"  # costs that move with revenue, saved each bad-quarter week
    LR = [LP_ROW + 1 + i for i in range(3)]

    def lows(a, e, shift, col):
        """Sum over the low periods of cut x the overlap of [a, e) with the period shifted by `shift` weeks."""
        terms = [f"N(Inputs!${col}${r})*MAX(0,MIN({e},{shift}+Inputs!$G${r})-MAX({a},{shift}+Inputs!$F${r}))" for r in LR]
        return "(" + "+".join(terms) + ")"
    first = 5
    sc = res.get("scenario") or {}
    nweeks = len(sc["base"]) if sc else 13
    m3 = sc.get("end_week") or sc.get("m3_week") or nweeks
    for w in range(1, nweeks + 1):
        r = first + w - 1
        cs.cell(row=r, column=1, value=w)
        cs.cell(row=r, column=2, value=f"={I('date')}+7*(A{r}-1)").number_format = "mmm d" if STYLE["dates"] in ("mdy", "ymd") else "dd mmm"
        a, e = f"(A{r}-1)", f"A{r}"
        cs.cell(row=r, column=3, value=f"={W}*(1-{lows(a, e, '0', 'D')})")
        cs.cell(row=r, column=4, value=f"=({S}+({W}-{S})*(1-{CUT}))*(MAX(0,{e}-MAX({a},{D}))-{lows(a, e, D, 'D')})")
        cs.cell(row=r, column=5, value=f"={I('costs_month_total')}*12/52*(1-{lows(a, e, '0', 'E')})")
        cs.cell(row=r, column=6, value=f"=IF(A{r}>={I('start_week')},N({I('role_monthly')})*12/52,0)+IF(A{r}={I('start_week')},N({I('one_off')}),0)")
        cs.cell(row=r, column=7, value=f"=SUMIF(Inputs!$D${KP_ROW + 1}:$D${KP_ROW + 10},A{r},Inputs!$B${KP_ROW + 1}:$B${KP_ROW + 10})")
        pH = "$B$2" if w == 1 else f"H{r - 1}"
        pI = "$B$2" if w == 1 else f"I{r - 1}"
        pM = "$B$2" if w == 1 else f"M{r - 1}"
        pP = "$B$2" if w == 1 else f"P{r - 1}"
        cs.cell(row=r, column=8, value=f"={pH}+C{r}-E{r}-F{r}-G{r}")
        cs.cell(row=r, column=9, value=f"={pI}+D{r}-E{r}+{REL}-F{r}-G{r}")
        cs.cell(row=r, column=10, value=f"={I('cover_target_months')}*({I('costs_month_total')}+IF(A{r}>={I('start_week')},N({I('role_monthly')}),0))")
        cs.cell(row=r, column=11, value=f'=IF(I{r}<J{r}-0.000001,"below","")')
        crate = f"N({I('contractor_rate')})"
        cs.cell(row=r, column=12, value=f"=IF(A{r}>={I('start_week')},{crate}*N({I('role_hours')}),0)")
        cs.cell(row=r, column=13, value=f"={pM}+D{r}-E{r}+{REL}-L{r}-G{r}")
        cs.cell(row=r, column=14, value=f"={I('cover_target_months')}*({I('costs_month_total')}+IF(A{r}>={I('start_week')},{crate}*N({I('role_hours')})*52/12,0))")
        lineh = f"MIN(N({I('role_hours')}),{I('contractor_line')})"
        cs.cell(row=r, column=15, value=f"=IF(A{r}>={I('start_week')},{crate}*{lineh},0)")
        cs.cell(row=r, column=16, value=f"={pP}+D{r}-E{r}+{REL}-O{r}-G{r}")
        cs.cell(row=r, column=17, value=f"={I('cover_target_months')}*({I('costs_month_total')}+IF(A{r}>={I('start_week')},{crate}*{lineh}*52/12,0))")
        for col in range(3, 18):
            if col != 11:
                cs.cell(row=r, column=col).number_format = mfmt
    last = first + nweeks - 1
    m3r = first + m3 - 1
    cs.cell(row=last + 2, column=1, value="The hire's 13th week, usual").font = bold
    cs.cell(row=last + 2, column=2, value=f"=H{m3r}").number_format = mfmt
    cs.cell(row=last + 3, column=1, value="The hire's 13th week, bad quarter").font = bold
    cs.cell(row=last + 3, column=2, value=f"=I{m3r}").number_format = mfmt
    cs.cell(row=last + 4, column=1, value="First week below floor").font = bold
    cs.cell(row=last + 4, column=2, value=f'=IFERROR(MATCH("below",K{first}:K{last},0),"none")')
    cs.cell(row=last + 5, column=1, value="Most the bad quarter falls short of the floor").font = bold
    for w in range(nweeks):
        r = first + w
        cs.cell(row=r, column=19, value=f"=J{r}-I{r}").number_format = mfmt
    cs.cell(row=4, column=19, value="Gap to floor, bad quarter").font = bold
    cs.cell(row=last + 5, column=2, value=f"=MAX(0,MAX(S{first}:S{last}))").number_format = mfmt
    for j in range(1, 20):
        cs.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 14
    cs.row_dimensions[4].height = 45

    cd = wb.create_sheet("Conditions")
    cd["A1"] = "The six conditions"
    cd["A1"].font = Font(bold=True, size=13)
    for j, h in enumerate(("#", "Condition", "Result", "Test"), 1):
        cd.cell(row=3, column=j, value=h).font = bold
    c5 = next(c for c in res["conditions"] if c["n"] == 5)
    H = lambda i: f"Inputs!$B${HIST_ROW + i}"  # noqa: E731
    hr = f"{H(0)}:{H(HIST_N - 1)}"
    last6 = f"{H(HIST_N - 6)}:{H(HIST_N - 1)}"
    prev6 = f"{H(HIST_N - 12)}:{H(HIST_N - 7)}"
    year6 = f"{H(HIST_N - 18)}:{H(HIST_N - 13)}"
    if len(gm) < 2 and rough:
        c2 = (f"=IF(COUNT({I('gm1')},{I('gm2')})<2,\"Unknown\",IF({I('gm1')}>={I('margin_floor')},\"Pass\","
              f"IF({I('gm2')}<{I('margin_floor')},\"Fail\",\"Unknown\")))")
        c2t = "Your rough range: Pass if the low end clears the floor, Fail if the high end is below it, Unknown in between"
    else:
        c2 = f"=IF(COUNT({I('gm1')},{I('gm2')})<2,\"Unknown\",IF(AND({I('gm1')}>={I('margin_floor')},{I('gm2')}>={I('margin_floor')}),\"Pass\",\"Fail\"))"
        c2t = "Both of the last two months at or above the floor"
    rows = [
        (1, "Cash holds in a bad quarter", f"=IF(N({I('role_monthly')})=0,\"Unknown\",IF(COUNTIF('Cash by week'!K{first}:K{last},\"below\")=0,\"Pass\",\"Fail\"))",
         "Bad-quarter bank at or above the floor in every week of the window"),
        (2, "Margin above your floor for two months", c2, c2t),
        (3, "Six months of the role covered",
         f"=IF(OR(N({I('role_monthly')})=0,AND({I('contracted_6')}=\"\",{I('recurring_month')}=\"\",{I('repeat_month')}=\"\")),\"Unknown\","
         f"IF(IF({I('contracted_6')}<>\"\",{I('contracted_6')},N({I('recurring_month')})*6)+N({I('repeat_month')})*6+N({I('pipeline_weighted')})>=6*{I('role_monthly')},\"Pass\",\"Fail\"))",
         "Recurring, contracted or repeat revenue for six months plus weighted pipeline, against six months of the role"),
        (4, "Your worst month is rising",
         f"=IF(COUNT({hr})>=18,IF(MIN({last6})>MIN({year6}),\"Pass\",\"Fail\"),IF(COUNT({hr})>=12,IF(MIN({last6})>MIN({prev6}),\"Pass\",\"Unknown\"),\"Unknown\"))",
         "At least 18 months: worst of the last six above the worst of the same six a year before. 12 to 17: above the six before, else Unknown (different seasons)"),
        (5, "A real role, and a use for your freed hours", c5["status"], "Judged in the report: named tasks, written down, and a use for the hours"),
        (6, "You are near full, and a price rise came first", f"=IF(COUNT({I('cap_now')},{I('cap_max')})<2,\"Unknown\",IF(AND({I('cap_now')}/{I('cap_max')}>={I('near_ceiling')},LOWER({I('price_rise')})=\"yes\"),\"Pass\",\"Fail\"))",
         "At or above the near-full share, and a price rise for new clients tested"),
    ]
    for i, (n, name, f, test) in enumerate(rows, 4):
        cd.cell(row=i, column=1, value=n)
        cd.cell(row=i, column=2, value=name)
        cd.cell(row=i, column=3, value=f)
        cd.cell(row=i, column=4, value=test)
    cd["A11"] = "Verdict"
    cd["A11"].font = bold
    no_rule = "No" if res["verdict"] == "No" and not res["verdict_reason"].startswith("Not as an employee") else ""
    cd["C11"] = (f"=IF(AND(LOWER({I('employment')})<>\"contractor\",N({I('role_hours')})>0,N({I('role_hours')})<{I('contractor_line')}),\"No\","
                 + (f"\"{no_rule}\"" if no_rule else f"IF(AND(COUNTIF(C4:C9,\"Pass\")=6,LOWER({I('cost_confirmed')})=\"yes\"),\"Ready\",\"Not yet\")") + ")")
    cd["C11"].font = Font(bold=True, size=13)
    cd["D11"] = ("Ready needs all six and the hire's fully loaded cost from your accountant or payroll provider. "
                 "No: an employee for under the contractor line, or a business already spending above its income.")
    cd.column_dimensions["B"].width = 46
    cd.column_dimensions["C"].width = 12
    cd.column_dimensions["D"].width = 90

    ct = wb.create_sheet("Contractor")
    ct["A1"] = "The contractor alternative"
    ct["A1"].font = Font(bold=True, size=13)
    ct["A3"] = "Hours a week, full role"
    ct["B3"] = f"=N({I('role_hours')})"
    ct["A4"] = "Monthly cost, full role"
    ct["B4"] = f"=N({I('contractor_rate')})*B3*52/12"
    ct["A5"] = "Passes the cash test?"
    ct["B5"] = (f"=IF(N({I('contractor_rate')})=0,\"[NEED: rate]\",IF(SUMPRODUCT(--('Cash by week'!M{first}:M{last}<'Cash by week'!N{first}:N{last}-0.000001))=0,\"Pass\",\"Fail\"))")
    ct["A7"] = "Hours a week, at the contractor line"
    ct["B7"] = f"=MIN(N({I('role_hours')}),{I('contractor_line')})"
    ct["A8"] = "Monthly cost, at the line"
    ct["B8"] = f"=N({I('contractor_rate')})*B7*52/12"
    ct["A9"] = "Passes the cash test?"
    ct["B9"] = (f"=IF(N({I('contractor_rate')})=0,\"[NEED: rate]\",IF(SUMPRODUCT(--('Cash by week'!P{first}:P{last}<'Cash by week'!Q{first}:Q{last}-0.000001))=0,\"Pass\",\"Fail\"))")
    for c in ("B4", "B8"):
        ct[c].number_format = mfmt
    ct.column_dimensions["A"].width = 40
    ct.column_dimensions["B"].width = 16
    wb.properties.title = f"Hiring readiness: {role['title']}"
    wb.properties.description = CREDIT
    try:
        from openpyxl.workbook.properties import CalcProperties
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
    except Exception:  # noqa: BLE001
        pass
    wb.save(path)


def write_csv(folder, stem, res):
    p = folder / f"{stem}_cash-workings.csv"
    sc = res.get("scenario")
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Week", "Bank, usual quarter", "Bank, bad quarter", "Floor"])
        if sc:
            for bw, sw in zip(sc["base"], sc["stress"]):
                fl = sc["floor"] if bw["week"] >= res["role"]["start_week"] else sc["floor_before_start"]
                w.writerow([bw["week"], round(bw["balance"], 2), round(sw["balance"], 2), round(fl, 2)])
    return p


# ------------------------------------------------------------------ pack only
def pack_only(a):
    from role_pack import check as rp_check
    if not a.role_pack or not a.inputs:
        raise InputError("--pack-only needs --role-pack and --inputs")
    for f in (a.role_pack, a.inputs):
        if not Path(f).exists():
            raise InputError(f"there is no file at {f}")
    rp = json.loads(Path(a.role_pack).read_text(encoding="utf-8"))
    inp = json.loads(Path(a.inputs).read_text(encoding="utf-8"))
    probs, _w = rp_check(rp, inp)
    if probs:
        raise InputError("the role pack fails its check (" + "; ".join(probs[:3]) + "); run role_pack.py check and fix it")
    role = inp.get("role") or {}
    apply_style(inp)
    today = date.fromisoformat(inp["date"]) if inp.get("date") else date.today()
    res = {"role": {"title": role.get("title") or "the role"}}
    L = [f"# Role pack: {res['role']['title']}",
         f"_{(inp.get('business') + '. ') if inp.get('business') and inp.get('business') != 'Your business' else ''}{long_date(today)}. A planning tool, not advice._",
         "**Readiness not checked, at your request.** Whether the cash holds with this cost added in a bad quarter "
         "is not tested here; run the full check before you commit."]
    L += role_pack_md(rp, res)
    L.append("## Before you act")
    L.append(f"This is a planning tool, not legal or employment advice. Contracts, payroll and whether this person would be an employee or a contractor: {PROS}.")
    md = "\n".join(L) + "\n"
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    p = out / f"role-pack_{slug(res['role']['title'])}_{today.isoformat()}.docx"
    write_docx(p, md, f"Role pack: {res['role']['title']}", FOOTER, STYLE["paper"], lang_of(inp.get("locale")), credit=CREDIT)
    if a.md_out:
        Path(a.md_out).write_text(md, encoding="utf-8")
    print(f"WROTE: {p}")
    return 0


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("results", nargs="?")
    ap.add_argument("--pack-only", action="store_true", help="write the role pack alone (the founder skipped the check)")
    ap.add_argument("--inputs", help="with --pack-only: the inputs file (business, role, tasks)")
    ap.add_argument("--out", required=True)
    ap.add_argument("--note")
    ap.add_argument("--lead", help="lines the founder wants first (copied from the conditions output); under the verdict unless --lead-above")
    ap.add_argument("--lead-above", action="store_true", help="put the lead lines above the verdict, when the founder asks for that")
    ap.add_argument("--role-pack")
    ap.add_argument("--md-out")
    ap.add_argument("--format", default="auto", choices=["auto", "xlsx", "csv"])
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args()
    if a.pack_only:
        return pack_only(a)
    if not a.results:
        raise InputError("give the results file from conditions.py")
    rp_path = Path(a.results)
    if not rp_path.exists():
        raise InputError(f"there is no results file at {a.results}; run conditions.py first")
    res = json.loads(rp_path.read_text(encoding="utf-8"))
    if not isinstance(res, dict) or "conditions" not in res:
        raise InputError("that file is not a results file from conditions.py")
    apply_style(res.get("inputs"))
    rp = None
    if a.role_pack:
        p = Path(a.role_pack)
        if not p.exists():
            raise InputError(f"there is no role pack file at {a.role_pack}")
        rp = json.loads(p.read_text(encoding="utf-8"))
        from role_pack import check as rp_check
        probs, _w = rp_check(rp, {"tasks": res["inputs"]["tasks"]})
        if probs:
            raise InputError("the role pack fails its check (" + "; ".join(probs[:3]) + "); run role_pack.py check and fix it")
    note = None
    if a.note:
        p = Path(a.note)
        if not p.exists():
            raise InputError(f"there is no note file at {a.note}")
        note = p.read_text(encoding="utf-8").strip()
        style = res["inputs"].get("decimal_style")
        allowed = [(v, k) for v, k, _e in tokens(report_md(res, None, rp), style)]
        allowed += [(v, k) for v, k, _e in tokens(json.dumps(res.get("scenario") or {}) + json.dumps(res["inputs"], default=str))]
        allowed += [(v, "money") for v in money_values(res)]
        probs = check_note(note, allowed, style)
        if probs:
            raise InputError("the note quotes figures or claims the results do not hold (" + "; ".join(probs) + "). Copy figures from the conditions output")
    lead = None
    if a.lead:
        p = Path(a.lead)
        if not p.exists():
            raise InputError(f"there is no lead file at {a.lead}")
        lead = p.read_text(encoding="utf-8").strip()
        style = res["inputs"].get("decimal_style")
        allowed = [(v, k) for v, k, _e in tokens(report_md(res, None, rp), style)]
        allowed += [(v, k) for v, k, _e in tokens(json.dumps(res.get("scenario") or {}) + json.dumps(res["inputs"], default=str))]
        allowed += [(v, "money") for v in money_values(res)]
        probs = check_note(lead, allowed, style)
        if probs:
            raise InputError("the lead lines quote figures or claims the results do not hold (" + "; ".join(probs) + "). Copy them from the conditions output")
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    stem = f"hiring-readiness_{slug(res['role']['title'])}_{res['date']}"
    fmt = a.format
    if fmt == "auto":
        try:
            import openpyxl  # noqa: F401
            fmt = "xlsx"
        except ImportError:
            fmt = "csv"
    md = report_md(res, note, rp, csv_fallback=(fmt == "csv"), lead=lead, lead_above=a.lead_above)
    written = []
    try:
        docx = out / f"{stem}.docx"
        write_docx(docx, md, f"Hiring readiness: {res['role']['title']}", FOOTER, STYLE["paper"], lang_of(res["inputs"].get("locale")), credit=CREDIT)
        written.append(docx)
        if fmt == "xlsx":
            xl = out / f"{stem}_cash-workings.xlsx"
            write_xlsx(xl, res)
            written.append(xl)
            if not a.no_recalc and not recalc(xl):
                print("NOTE: calculated values were not stored (LibreOffice not found); Excel, Numbers and Google Sheets calculate on open.")
        else:
            written.append(write_csv(out, stem, res))
            print("NOTE: openpyxl is not available, so the cash workings are a CSV file; the report carries every figure.")
        if a.md_out:
            Path(a.md_out).parent.mkdir(parents=True, exist_ok=True)
            Path(a.md_out).write_text(md, encoding="utf-8")
    except Exception:
        for p in written:
            p.unlink(missing_ok=True)
        raise
    for p in written:
        print(f"WROTE: {p}")
    print(f"VERDICT: {res['verdict']}")
    return 1 if res["needs"] else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except InputError as e:
        print(f"STOP: {e}. Nothing was written.")
        sys.exit(2)
    except Exception as e:  # noqa: BLE001
        print(f"STOP: the files could not be built ({type(e).__name__}: {e}). Nothing usable was written.")
        sys.exit(2)
