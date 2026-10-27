#!/usr/bin/env python3
"""Roll a 13-week cash forecast, week by week, in a normal and a bad quarter, against a safety line.

Usage:
  python3 scripts/forecast.py compute <inputs.json>
        print every figure the reading may use; exit 1 when there are [NEED] gaps
  python3 scripts/forecast.py build <inputs.json> --out <output folder> [--note <note.md>]
        [--previous <previous.json>] [--results <scratch>/results.json] [--format auto|xlsx|csv] [--no-recalc]
        write cash-forecast-<business>-<date>.xlsx (live formulas) and the one-page reading (.docx)
  python3 scripts/forecast.py template
        print an empty inputs file

<inputs.json> starts as the draft find_recurring.py writes; Claude adds the founder's answers.
Every week asserts closing = opening + money in less money out. Exit codes: 0 written, 1 written with
[NEED] gaps, 2 STOP (nothing usable written; read the line). Standard library only, plus openpyxl when
present for the workbook (CSV files otherwise). No network calls.
"""
import argparse
import json
import re
import statistics
import sys
import zipfile
from datetime import date, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cashlib import (NUM, SYMBOL_AFTER, SYMBOLS, InputError, add_months, from_iso, iso, long_date, money,  # noqa: E402
                     month_end, parse_count, parse_money, parse_pct, parse_when, set_account_currency, set_decimal,
                     set_number_style, set_style, short_date, style_from_samples, typed_samples, TYPED)
import locale_numbers as ln  # noqa: E402

SKILL_NAME = "13-Week Cash Forecast"
HORIZON_WEEKS = 13
# Working defaults. Merigrove method (references/method.md section 6); Max reviews before release.
DEFAULT_STRESS_CUT = 0.20        # receipts 20 per cent lower in a bad quarter
DEFAULT_STRESS_DELAY_DAYS = 15   # and 15 days later
DEFAULT_PLAN_DEFAULT = 0.10      # payment plan instalments that fail in a bad quarter, until the founder has history
RUNWAY_MIN_WEEKS = 8             # stress-case runway the method asks for
PROXIMITY_WEEKS = 4              # a breach this close means act this week and re-run weekly
DEFAULT_TERMS_DAYS = 30          # when invoices are listed: new work is paid this long after today
DEFAULT_PAYDAY = 28              # for a hire when no payroll day is known
CADENCES = {"weekly": 7, "fortnightly": 14, "four-weekly": 28, "monthly": None, "quarterly": None, "annual": None}
MONTHLY_FACTOR = {"weekly": 52 / 12, "fortnightly": 26 / 12, "four-weekly": 13 / 12, "monthly": 1.0,
                  "quarterly": 1 / 3, "annual": 1 / 12}
NOT_IN_FLOOR = {"own transfer", "tax set-aside", "linked cost"}
FREQUENT = {"weekly", "fortnightly", "four-weekly", "monthly"}  # known payments this frequent count in the floor
KNOWN_KINDS = {"tax", "payroll tax", "government", "other"}
_DF = {"v": None}  # the founder's date order for typed dates: True day first, False month first, None not known


def _when(v, today):
    return parse_when(v, today, _DF["v"])


# ------------------------------------------------------------------ reading the inputs
def pct(v, default, label):
    if v in (None, ""):
        return default
    x = parse_pct(v, label)
    if not 0 <= x < 0.9:
        raise InputError(f"{label} of {round(x * 100)}% is outside 0 to 90 per cent")
    return x


def amt(v, label, needs=None, allow_none=True):
    x = parse_money(v, label)
    if x is None:
        if not allow_none:
            raise InputError(f"{label} is missing")
        return None
    return abs(x)


def days_of(v, label, default=None):
    x = parse_count(v, label)
    return default if x is None else x


def clean_label(s, fallback):
    s = re.sub(r"[*?~]", " ", str(s or fallback))
    s = re.sub(r"\s+", " ", s).strip().lstrip("=+-@").strip()
    return (s or fallback)[:60]


def cadence_of(v):
    s = str(v or "monthly").lower().strip()
    s = {"week": "weekly", "every week": "weekly", "a week": "weekly", "pw": "weekly", "per week": "weekly",
         "biweekly": "fortnightly", "every two weeks": "fortnightly", "every 2 weeks": "fortnightly",
         "month": "monthly", "a month": "monthly", "pm": "monthly", "per month": "monthly", "every month": "monthly",
         "4-weekly": "four-weekly", "every 4 weeks": "four-weekly", "quarter": "quarterly", "every quarter": "quarterly",
         "year": "annual", "yearly": "annual", "annually": "annual", "a year": "annual"}.get(s, s)
    if s not in CADENCES:
        raise InputError(f"'{v}' is not a cadence I know; use weekly, fortnightly, four-weekly, monthly, quarterly or annual")
    return s


def validate_shape(raw):
    if not isinstance(raw, dict):
        raise InputError("the inputs file must be one JSON object (start from find_recurring.py or forecast.py template)")
    for k in ("recurring_out", "recurring_in", "invoices", "instalments", "one_offs"):
        if raw.get(k) is not None and not isinstance(raw.get(k), list):
            raise InputError(f"'{k}' must be a list")
    if raw.get("opening") is not None and not isinstance(raw["opening"], dict):
        raise InputError("'opening' must be an object with balance, date and source")


def locale_setup(raw, notes):
    """Currency, decimal style, the way figures are written back, date order and paper size, for this founder.
    Paper and date order come from the founder's country (locale_numbers.locale_defaults), never from spelling."""
    cur = (str(raw.get("currency") or "").strip().upper()) or None
    set_account_currency(cur)
    samples = typed_samples(raw)
    ex = raw.get("number_style") or {}
    given = str(raw.get("decimal_style") or "").lower()
    if given in ("dot", "comma"):
        dec, dec_src = given, "your answer"
    else:
        dec = ln.detect_decimal_style(samples)
        dec_src = "what you typed" if dec else None
        if not dec and ex.get("decimal") in (".", ","):
            dec, dec_src = ("comma" if ex["decimal"] == "," else "dot"), "your bank export"
    set_decimal(dec)
    g_typed, d_typed = style_from_samples(samples)
    group = g_typed or ex.get("group") or ","
    decimal = d_typed or {"dot": ".", "comma": ","}.get(dec) or ex.get("decimal") or "."
    set_number_style(group, decimal)
    NUM["after"] = bool(cur and (cur in SYMBOL_AFTER or any(
        re.search(r"\d\s*" + re.escape(SYMBOLS.get(cur, "\x00")) + r"\s*$", str(x)) for x in samples)))
    country = (str(raw.get("country") or "").strip().upper()) or None
    loc = ln.locale_defaults(country)
    order = str(raw.get("date_order") or "").lower()
    if not order and country and loc["date_order"] in ("MDY", "DMY"):
        order = "month-first" if loc["date_order"] == "MDY" else "day-first"
        notes.append(f"dates you typed are read {order.replace('-', ' ')}, as written in {country}")
    if not country:
        notes.append("the business's country is not set, so the one-page reading is laid out for A4; set country "
                     "(for example US, ZA, SG) for Letter or A4 as used there")
    return {"currency": cur, "decimal": dec, "decimal_src": dec_src, "group": NUM["group"], "decimal_mark": NUM["decimal"],
            "country": country, "paper": loc["paper"], "order": order}


def normalise(raw):
    validate_shape(raw)
    needs, notes = [], []
    lc = locale_setup(raw, notes)
    order = lc["order"]
    _DF["v"] = True if order.startswith("day") else (False if order.startswith("month") else None)
    set_style(month_first=_DF["v"] is False)
    today = parse_when(raw.get("today"), date.today(), True) if raw.get("today") else date.today()
    cur = lc["currency"]
    if not cur:
        needs.append("the account's currency is not confirmed, so figures show without a currency sign. Ask: which "
                     "currency is this account in?")
    op = raw.get("opening") or {}
    bal = parse_money(op.get("balance"), "the opening balance") if op.get("balance") not in (None, "") else None
    if bal is None:
        raise InputError("there is no opening balance. Ask the founder: 'What is in the business account today?' and "
                         "put it in opening.balance with today's date")
    odate = _when(op.get("date"), today) if op.get("date") else today
    if odate > today + timedelta(days=1):
        raise InputError(f"the opening balance is dated {long_date(odate)}, after today")
    start = odate + timedelta(days=1)
    end = start + timedelta(days=7 * HORIZON_WEEKS - 1)
    _YEARS["v"] = start.year != end.year
    _YEARS["date_first"] = str(raw.get("week_labels") or "").lower().replace("-", " ").startswith("date first")
    m = {"business": clean_label(raw.get("business"), "Your business"), "currency": cur, "today": today,
         "country": lc["country"], "paper": lc["paper"], "decimal_style": lc["decimal"],
         "number_style": {"group": lc["group"], "decimal": lc["decimal_mark"], "after": NUM.get("after", False)},
         "asked_about": [str(x) for x in (raw.get("asked_about") or []) if x],
         "moves_first": [str(x) for x in (raw.get("moves_first") or []) if x] if isinstance(raw.get("moves_first"), list)
         else ([str(raw["moves_first"])] if raw.get("moves_first") else []),
         "moves_on_top": bool(raw.get("moves_on_top")),
         "date_order": "day-first" if _DF["v"] else ("month-first" if _DF["v"] is False else None),
         "opening": round(bal, 2), "opening_date": odate, "opening_source": op.get("source") or "founder said",
         "start": start, "end": end, "export": raw.get("export") or {}}
    if (today - odate).days > 10:
        notes.append(f"the opening balance is from {long_date(odate)}, {(today - odate).days} days ago")
    st = raw.get("stress") or {}
    m["cut"] = pct(st.get("receipts_cut"), DEFAULT_STRESS_CUT, "the bad-quarter receipts cut")
    m["delay"] = days_of(st.get("delay_days"), "the bad-quarter delay in days", DEFAULT_STRESS_DELAY_DAYS)
    if not 0 <= m["delay"] <= 120:
        raise InputError(f"a bad-quarter delay of {m['delay']} days is outside 0 to 120 days")
    m["plan_default"] = pct(st.get("plan_default"), DEFAULT_PLAN_DEFAULT, "the payment plan default rate")
    m["stress_src"] = "your figures" if st and (st.get("receipts_cut") not in (None, "20%", 0.2, "20") or
                                                  str(st.get("delay_days")) not in ("15", "None")) else "default"

    def items(key):
        return [x for x in (raw.get(key) or []) if isinstance(x, dict)]

    m["out"] = []
    for i, r in enumerate(items("recurring_out")):
        if r.get("keep") is False:
            notes.append(f"'{r.get('label')}' left out: it looks like it has stopped")
            continue
        a = amt(r.get("amount"), f"regular payment '{r.get('label')}'")
        if a is None:
            needs.append(f"no amount for '{r.get('label')}'; it is left out until you give one")
            continue
        cad = cadence_of(r.get("cadence"))
        nxt = _when(r.get("next_date"), today) if r.get("next_date") else None
        if nxt is None:
            raise InputError(f"regular payment '{r.get('label')}' has no next date")
        linked = r.get("linked_to") or (r.get("revenue_linked") and "work already booked")
        m["out"].append({"label": clean_label(r.get("label"), f"Regular payment {i + 1}"), "match": r.get("match"),
                         "category": "linked cost" if linked else (r.get("category") or "other regular").lower(),
                         "amount": a, "cadence": cad,
                         "next": nxt, "until": _when(r["until"], today) if r.get("until") else None,
                         "source": (r.get("source") or "founder said") + (f"; for {linked}, so never held" if linked else "")})
    ev = raw.get("everyday_spending_week") or {}
    m["everyday"] = amt(ev.get("amount"), "everyday spending") or 0.0
    m["everyday_src"] = ev.get("source") or "founder said"
    ow = raw.get("owner_pay_irregular_week") or {}
    m["owner_irreg"] = amt(ow.get("amount"), "irregular owner pay") or 0.0
    lates = {str(k).lower().strip(): days_of(v, f"days late for {k}", 0) for k, v in (raw.get("client_late_days") or {}).items()}
    m["lates"] = lates

    def late_for(client, given):
        if given not in (None, ""):
            return days_of(given, f"days late for {client}", 0), "your figure"
        c = str(client or "").lower().strip()
        for k, v in lates.items():
            if k and (k in c or c in k):
                return v, "saved habit"
        return 0, "default: on the due date"
    m["late_for"] = late_for
    m["ins"] = []
    for i, r in enumerate(items("recurring_in")):
        a = amt(r.get("amount"), f"regular receipt '{r.get('label')}'")
        if a is None:
            needs.append(f"no amount for regular receipt '{r.get('label')}'")
            continue
        nxt = _when(r.get("next_date"), today) if r.get("next_date") else None
        if nxt is None:
            raise InputError(f"regular receipt '{r.get('label')}' has no next date")
        late, lsrc = late_for(r.get("client") or r.get("label"), r.get("late_days"))
        m["ins"].append({"label": clean_label(r.get("label"), f"Regular client {i + 1}"), "match": r.get("match"),
                         "client": r.get("client") or r.get("label"),
                         "amount": a, "cadence": cadence_of(r.get("cadence")), "next": nxt, "late": late, "late_src": lsrc,
                         "until": _when(r["until"], today) if r.get("until") else None,
                         "source": r.get("source") or "founder said"})
    m["invoices"] = []
    for i, r in enumerate(items("invoices")):
        a = amt(r.get("amount"), f"invoice {i + 1}")
        if a is None:
            needs.append(f"invoice {i + 1} ({r.get('client')}) has no amount; left out")
            continue
        if r.get("due"):
            due = _when(r["due"], today)
        elif r.get("issued"):
            due = _when(r["issued"], today) + timedelta(days=days_of(r.get("terms_days"), "payment terms in days", DEFAULT_TERMS_DAYS))
        else:
            needs.append(f"invoice {i + 1} ({r.get('client')}) has no due date; placed {DEFAULT_TERMS_DAYS} days from today")
            due = today + timedelta(days=DEFAULT_TERMS_DAYS)
        late, lsrc = late_for(r.get("client"), r.get("late_days"))
        m["invoices"].append({"client": clean_label(r.get("client"), f"Client {i + 1}"), "amount": a, "due": due,
                              "late": late, "late_src": lsrc, "ref": r.get("ref") or ""})
    pat = raw.get("receipts_pattern_week") or {}
    m["pattern"] = amt(pat.get("amount"), "the usual receipts pattern") or 0.0
    if pat.get("use") is False:
        m["pattern"] = 0.0
    m["pattern_src"] = pat.get("source") or "founder said"
    m["pattern_months"] = pat.get("months_covered")
    # a client whose invoices are listed comes out of the usual pattern and out of the regular clients, so the
    # same money is never counted twice
    inv_names = [i["client"] for i in m["invoices"]]
    saved_inv = [str(x) for x in (pat.get("exclude") or []) if str(x).strip()]
    # a first-read what-if: the usual pattern without one payer, to show what happens if that client pays nothing more
    what_if = [str(x) for x in ([raw["without_payer"]] if isinstance(raw.get("without_payer"), str) else (raw.get("without_payer") or []))
               if str(x).strip()]
    m["pattern_removed"] = []
    weeks = float(pat.get("weeks") or 0) or max(1.0, float((raw.get("export") or {}).get("days") or 0) / 7)
    if m["pattern"] and pat.get("payers"):
        cut = 0.0
        for py in pat["payers"]:
            if not isinstance(py, dict):
                continue
            who = next((n for n in inv_names + saved_inv + what_if if same_client(n, py)), None)
            if who:
                cut += float(py.get("amount") or 0)
                if who not in m["pattern_removed"]:
                    m["pattern_removed"].append(who)
        if cut:
            before = m["pattern"]
            m["pattern"] = round(max(0.0, before - cut / weeks), 2)
            notes.append(f"usual pattern now {money(m['pattern'], cur)} a week, down from {money(before, cur)}, without "
                         f"{' and '.join(m['pattern_removed'])}: "
                         + ("a what-if for the first read: no more money from them in 13 weeks" if what_if and not inv_names
                            else "their money in comes from the invoices listed"))
    for n in saved_inv:
        if not any(same_client(n, {"payer": i}) for i in inv_names):
            needs.append(f"no open invoices are listed for {n}, whose receipts are left out of the usual pattern because "
                         "you list their invoices. List the open ones, or say they owe nothing")
    keep_ins = []
    for r in m["ins"]:
        if any(same_client(n, {"payer": r["client"], "label": r["label"]}) for n in inv_names):
            notes.append(f"regular receipts from {r['label']} are replaced by the invoices you listed for them")
            continue
        keep_ins.append(r)
    m["ins"] = keep_ins
    m["months_off"] = months_of(pat.get("months_off"))
    m["month_factor"] = {}
    for k, v in (pat.get("month_factor") or {}).items():
        mm_ = months_of([k])
        if mm_:
            m["month_factor"][mm_[0]] = float(parse_pct(v, f"the pattern for {k}") if isinstance(v, str) else v)
    terms = raw.get("payment_terms_days")
    m["terms"] = days_of(terms, "payment terms in days", DEFAULT_TERMS_DAYS)
    m["terms_src"] = "your figure" if terms not in (None, "") else "default"
    m["pattern_from"] = start + timedelta(days=m["terms"]) if m["invoices"] else start
    m["plans"] = []
    for i, r in enumerate(items("instalments")):
        dates = [_when(x, today) for x in (r.get("dates") or [])]
        each = amt(r.get("amount") or r.get("amount_each"), f"payment plan {i + 1}")
        if not dates and r.get("first") and r.get("count"):
            first = _when(r["first"], today)
            dates = [add_months(first, k) for k in range(days_of(r["count"], "the number of instalments", 0))]
        if each is None or not dates:
            needs.append(f"payment plan {i + 1} ({r.get('client')}) needs an amount and dates")
            continue
        m["plans"].append({"client": clean_label(r.get("client"), f"Plan {i + 1}"), "amount": each, "dates": dates})
    m["one_offs"] = []
    for i, r in enumerate(items("one_offs")):
        a = amt(r.get("amount"), f"one-off '{r.get('label')}'")
        if a is None or not r.get("date"):
            needs.append(f"one-off '{r.get('label')}' needs an amount and a date")
            continue
        lab = clean_label(r.get("label"), f"One-off {i + 1}")
        linked = r.get("linked_to") or next((n for n in inv_names if same_client(n, {"payer": lab})), None)
        m["one_offs"].append({"label": lab, "amount": a,
                              "date": _when(r["date"], today),
                              "direction": "in" if str(r.get("direction", "out")).lower().startswith("in") else "out",
                              "certain": bool(r.get("certain")), "linked": linked})
    known_normalise(raw, m, needs, today)
    provisional_tax(raw, m, notes, today)
    m["commit"] = commit_normalise(raw.get("commitment"), m, today, needs)
    m["linked_cost"] = None
    if m["commit"] and m["commit"]["linked"]:
        m["linked_cost"], m["commit"] = m["commit"], None
        notes.append(f"'{m['linked_cost']['label']}' is tied to work already booked, so it is in the forecast as a cost, "
                     "not tested as a choice")
    floor_normalise(raw, m, needs)
    m["history_daily"] = raw.get("history_daily") or {}
    plausibility(raw, m)
    m["tax_seen"] = [t for t in ((raw.get("seen") or {}).get("tax_like") or []) if isinstance(t, dict)]
    m["sources"] = raw.get("sources") or {}
    return m, needs, notes


GENERIC_NAME = {"holdings", "group", "services", "solutions", "company", "corporate", "international", "global",
                "limited", "partners", "consulting", "studio", "agency", "trust", "bank", "credit", "payment", "payments",
                "transfer", "invoice", "invoices", "deposit", "balance", "wedding", "function", "event", "events",
                "lunch", "lunches", "year", "suppliers", "supplier", "client", "clients", "health", "insurance"}


def name_words(s):
    return {w for w in re.findall(r"[a-z]{4,}", str(s or "").lower()) if w not in GENERIC_NAME}


def same_client(name, payer):
    """Whether a client the founder named is this payer in the export: a distinctive word in common."""
    a = name_words(name)
    b = name_words(payer.get("payer")) | name_words(payer.get("label"))
    return bool(a & b)


MONTH_NAMES = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
               "november", "december"]


def months_of(v):
    """Month numbers from what the founder said: ['jan', 'Feb'], ['January'], [1, 2], 'jan-feb'."""
    out = []
    items = v if isinstance(v, list) else re.split(r"[,;/ ]+|\band\b", str(v or ""))
    for x in items:
        x = str(x).strip().lower()
        rng = re.fullmatch(r"([a-z]{3,9})\s*(?:-|to|through)\s*([a-z]{3,9})", x)
        if rng:
            a, b = months_of([rng.group(1)]), months_of([rng.group(2)])
            if a and b:
                k = a[0]
                while True:
                    out.append(k)
                    if k == b[0]:
                        break
                    k = k % 12 + 1
            continue
        if x.isdigit() and 1 <= int(x) <= 12:
            out.append(int(x))
        elif len(x) >= 3:
            hit = [i + 1 for i, n in enumerate(MONTH_NAMES) if n.startswith(x[:3])]
            if hit:
                out.append(hit[0])
    return sorted(set(out))


def plausibility(raw, m):
    """Stop with one plain question when a headline figure the founder typed cannot be right (a misread decimal,
    a missing zero). A figure the founder has confirmed is listed in 'confirmed' and passes."""
    ok_ = {str(x).lower().strip() for x in (raw.get("confirmed") or [])}
    cur = m["currency"]
    hist = m.get("history_daily") or {}
    ins_ = sorted(v[0] for v in hist.values() if isinstance(v, list) and v and v[0] > 0)
    typical_in = statistics.median(ins_) if ins_ else None
    biggest_in = max(ins_) if ins_ else None
    mo = m.get("monthly_out") or 0

    def stop(label, fig, why):
        if label.lower().strip() in ok_ or money(fig, cur).lower() in ok_:
            return
        raise InputError(f"{label} reads as {money(fig, cur, pence=True)}, {why}. Ask the founder once: 'Is "
                         f"{money(fig, cur, pence=True)} right for {label}?' If they confirm it, add '{label}' to "
                         "confirmed in the inputs")
    for i in m["invoices"]:
        lab = f"the invoice from {i['client']}"
        if i["amount"] < 1:
            stop(lab, i["amount"], "less than one unit of the currency")
        if typical_in and i["amount"] < typical_in / 100:
            stop(lab, i["amount"], f"a hundredth of a typical day's receipts in your export ({money(typical_in, cur)})")
        if biggest_in and i["amount"] > 50 * biggest_in and i["amount"] > 10 * abs(m["opening"] or 1):
            stop(lab, i["amount"], f"fifty times the largest day's receipts in your export ({money(biggest_in, cur)})")
    for k in m.get("known", []):
        if k["amount"] is None:
            continue
        if k["amount"] < 1:
            stop(f"'{k['label']}'", k["amount"], "less than one unit of the currency")
        if mo and k["amount"] > 12 * mo:
            stop(f"'{k['label']}'", k["amount"], f"more than a year of your outgoings ({money(mo, cur)} a month)")
    for o in m["one_offs"]:
        if o["amount"] < 1:
            stop(f"'{o['label']}'", o["amount"], "less than one unit of the currency")
        if mo and o["amount"] > 24 * mo and o["amount"] > 10 * abs(m["opening"] or 1):
            stop(f"'{o['label']}'", o["amount"], f"more than two years of your outgoings ({money(mo, cur)} a month)")
    c = m.get("commit") or m.get("linked_cost")
    if c:
        cost = c["monthly"] or c["one_off"]
        if cost and cost < 1:
            stop(f"the cost of '{c['label']}'", cost, "less than one unit of the currency")
        if mo and c["monthly"] > 10 * mo:
            stop(f"the monthly cost of '{c['label']}'", c["monthly"], f"ten times your current outgoings ({money(mo, cur)} a month)")
    if m["floor_basis"] == "the figure you set":
        if m["floor"] < 1:
            stop("the safety line", m["floor"], "less than one unit of the currency")
        if mo and m["floor"] > 24 * mo:
            stop("the safety line", m["floor"], f"more than two years of your outgoings ({money(mo, cur)} a month)")


def known_normalise(raw, m, needs, today):
    """Tax, payroll taxes and other government or fixed payments the founder lists, with the amount and the date
    or rhythm they give (from their accountant, bookkeeper or payroll provider). Never derived: a missing amount
    is placed at zero and marked [NEED]; a missing date is not placed and marked [NEED]."""
    m["known"] = []
    for i, r in enumerate(x for x in (raw.get("known_outgoings") or []) if isinstance(x, dict)):
        label = clean_label(r.get("label"), f"Known payment {i + 1}")
        kind = str(r.get("kind") or "other").lower()
        kind = kind if kind in KNOWN_KINDS else "other"
        a = amt(r.get("amount"), f"known payment '{label}'")
        who = str(r.get("source") or "you").strip()
        src = "founder supplied" + ("" if who.lower() in ("you", "") else f": {who}")
        dates, cad, nxt = [], None, None
        if r.get("dates"):
            dates = [_when(x, today) for x in r["dates"]]
        elif r.get("date"):
            dates = [_when(r["date"], today)]
        elif r.get("cadence") and r.get("next_date"):
            cad = cadence_of(r["cadence"])
            nxt = _when(r["next_date"], today)
        if not dates and not nxt:
            needs.append(f"'{label}' has no date or rhythm, so it is not in the forecast. Ask the founder (or their "
                         f"{'payroll provider' if kind == 'payroll tax' else 'accountant'}) when it is due")
            continue
        if a is None:
            needs.append(f"'{label}' has no amount: placed at zero until you give the figure from your "
                         f"{'payroll provider' if kind == 'payroll tax' else 'accountant'}")
        m["known"].append({"label": label, "kind": kind, "amount": a, "dates": dates, "cadence": cad, "next": nxt,
                           "until": _when(r["until"], today) if r.get("until") else None, "source": src, "who": who})


def provisional_tax(raw, m, notes, today):
    """First read only: repeat a tax-like payment the export shows at the amount the bank last showed, on the
    rhythm it showed, so the first answer is not kinder than the bank. Never computed, never in the files: the
    build refuses an inputs file that asks for it. The founder's own figures replace it."""
    want = [str(x).lower() for x in (raw.get("provisional_tax") or []) if str(x).strip()]
    if not want:
        return
    m["provisional"] = True
    for t in ((raw.get("seen") or {}).get("tax_like") or []):
        if not isinstance(t, dict) or not any(w in str(t.get("label", "")).lower() or str(t.get("label", "")).lower() in w
                                              for w in want):
            continue
        pays = sorted((from_iso(p["date"]), float(p["amount"])) for p in (t.get("payments") or []) if p.get("date"))
        if len(pays) < 2:
            continue
        rhythm = str(t.get("rhythm_seen") or "")
        last_amt = float(t.get("last_amount") or pays[-1][1])
        label = f"{clean_label(t.get('label'), 'Tax-like payment')} (as your bank last showed it)"
        src = "your bank's last payment, provisional until you give the figure"
        if rhythm.startswith("twice a month"):
            for d, _a in pays[-2:]:
                m["known"].append({"label": label, "kind": "tax", "amount": last_amt, "dates": [], "cadence": "monthly",
                                   "next": add_months(d, 1), "until": None, "source": src, "who": "your bank"})
        elif rhythm.startswith("monthly") or rhythm.startswith("quarterly"):
            cad = "monthly" if rhythm.startswith("monthly") else "quarterly"
            m["known"].append({"label": label, "kind": "tax", "amount": last_amt, "dates": [], "cadence": cad,
                               "next": add_months(pays[-1][0], 1 if cad == "monthly" else 3), "until": None,
                               "source": src, "who": "your bank"})
        else:
            continue
        notes.append(f"PROVISIONAL: '{t.get('label')}' repeated at {money(last_amt, m['currency'])} on the rhythm your bank "
                     "showed, for the first read only; replace it with the founder's figure before building")


def commit_normalise(c, m, today, needs):
    if not c:
        return None
    kind = "spend" if str(c.get("kind") or "").lower().startswith("spend") else "hire"
    start = _when(c.get("start"), today) if c.get("start") else None
    if start is None:
        needs.append("the commitment has no start date; not tested")
        return None
    label = clean_label(c.get("label"), "New hire" if kind == "hire" else "Planned spend")
    monthly = amt(c.get("monthly_cost"), "commitment monthly cost")
    one_off = amt(c.get("one_off"), "commitment one-off cost") or 0.0
    if kind == "hire" and monthly is None:
        needs.append(f"'{label}' has no monthly cost (pay plus employer costs, from you or your accountant); not tested")
        return None
    if kind == "spend" and monthly is None and not one_off:
        one_off = amt(c.get("amount"), "planned spend") or 0.0
        if not one_off:
            needs.append(f"'{label}' has no amount; not tested")
            return None
    payday = days_of(c.get("payday"), "the payday", 0) or None
    payday_src = "your figure" if payday else None
    if not payday:
        days = [o["next"].day for o in m["out"] if o["category"] == "payroll"]
        payday = days[0] if days else DEFAULT_PAYDAY
        payday_src = "your current payroll day" if days else "default: the 28th"
    end = _when(c["end"], today) if c.get("end") else None
    return {"kind": kind, "label": label, "monthly": monthly or 0.0, "one_off": one_off, "start": start, "end": end,
            "payday": payday, "payday_src": payday_src, "on_costs_src": c.get("source") or "founder said",
            "linked": bool(c.get("revenue_linked"))}


def first_payday(start, payday):
    try:
        d = date(start.year, start.month, payday)
    except ValueError:
        d = month_end(start.year, start.month)
    if d < start:
        nm = add_months(date(start.year, start.month, 1), 1)
        try:
            d = date(nm.year, nm.month, payday)
        except ValueError:
            d = month_end(nm.year, nm.month)
    return d


def floor_normalise(raw, m, needs):
    f = raw.get("floor") or {}
    mode = str(f.get("mode") or "one_month").lower()
    val = f.get("value")
    mo, parts = monthly_outgoings(m)
    m["monthly_out"], m["monthly_parts"] = mo, parts
    if isinstance(val, str) and re.search(r"week|month", val, re.I):
        n = ln.parse_number(re.sub(r"(?i)\s*(weeks?|months?)(\s+of\s+\w+)?\s*$", "", val) or "0", TYPED["decimal"])[0]
        mode, val = ("weeks", n) if "week" in val.lower() else ("weeks", n * 52 / 12)
    if mode == "figure" or (mode not in ("one_month", "weeks") and val not in (None, "")):
        x = amt(val, "the safety line")
        if x is None:
            raise InputError("the safety line was set as a figure but no figure was given")
        m["floor"], m["floor_basis"] = round(x, 2), "the figure you set"
    elif mode == "weeks":
        n = float(val) if isinstance(val, (int, float)) else float(ln.parse_number(str(val), TYPED["decimal"])[0] if val not in (None, "") else 0)
        if n <= 0:
            raise InputError("the safety line in weeks needs a number of weeks")
        m["floor"] = round(mo * 12 / 52 * n, 2)
        m["floor_basis"] = f"{n:g} weeks of outgoings"
    else:
        m["floor"], m["floor_basis"] = round(mo, 2), "one month of outgoings"
    if m["floor_basis"] != "the figure you set" and (m.get("linked_cost") or any(o["category"] == "linked cost" for o in m["out"])):
        m["floor_basis"] += ", leaving out costs tied to booked work"
    if m["floor"] <= 0:
        needs.append("the safety line came to zero because no regular outgoings are known; set a figure")


def monthly_outgoings(m):
    parts = {}
    for o in m["out"]:
        if o["category"] in NOT_IN_FLOOR:
            continue
        key = {"payroll": "pay", "owner pay": "pay", "pension": "pay"}.get(o["category"], "regular bills")
        parts[key] = parts.get(key, 0) + o["amount"] * MONTHLY_FACTOR[o["cadence"]]
    for k in m.get("known", []):
        if k["cadence"] in FREQUENT and k["amount"]:
            parts["known payments you listed"] = parts.get("known payments you listed", 0) + k["amount"] * MONTHLY_FACTOR[k["cadence"]]
    if m.get("everyday"):
        parts["everyday spending"] = m["everyday"] * 52 / 12
    if m.get("owner_irreg"):
        parts["owner pay, irregular"] = m["owner_irreg"] * 52 / 12
    order = ["pay", "known payments you listed", "regular bills", "everyday spending", "owner pay, irregular"]
    items = [(k, round(parts[k], 2)) for k in order if k in parts]
    return round(sum(v for _k, v in items), 2), items


def monthly_outgoings_from_raw(raw):
    """Used by find_recurring.py to preview the floor on the draft."""
    r = dict(raw)
    r.setdefault("opening", {"balance": 0, "date": raw.get("today")})
    if r["opening"].get("balance") in (None, ""):
        r["opening"] = dict(r["opening"], balance=0)
    mm, _n, _x = normalise(r)
    return mm["monthly_out"], mm["monthly_parts"]


# ------------------------------------------------------------------ the schedule
def occurrences(first, cadence, end, until=None):
    out, k, d = [], 0, first
    while d <= end and (until is None or d <= until):
        out.append(d)
        k += 1
        if CADENCES.get(cadence):
            d = first + timedelta(days=CADENCES[cadence] * k)
        else:
            d = add_months(first, {"monthly": 1, "quarterly": 3, "annual": 12}[cadence] * k)
        if k > 200:
            break
    return out


def schedule(m):
    """Every receipt and payment in the 13 weeks, with its base and bad-quarter amount and date."""
    start, end = m["start"], m["end"]
    rec, pay, taxes = [], [], []

    def r_add(line, frm, kind, amount, due, late, source):
        base_date = max(start, due + timedelta(days=late))
        if kind == "payment plan":
            s_amt, s_date = amount * (1 - m["plan_default"]), base_date
        elif kind == "certain":
            s_amt, s_date = amount, base_date
        else:
            s_amt, s_date = amount * (1 - m["cut"]), base_date + timedelta(days=m["delay"])
        rec.append({"line": line, "from": frm, "kind": kind, "amount": round(amount, 2), "due": due, "late": late,
                    "date": base_date, "s_amount": round(s_amt, 2), "s_date": s_date, "source": source})
    for r in m["ins"]:
        for d in occurrences(r["next"], r["cadence"], end, r["until"]):
            r_add(r["label"], r["client"], "regular client", r["amount"], d, r["late"], r["source"])
    for inv in m["invoices"]:
        r_add(f"Invoices: {inv['client']}", inv["client"], "invoice", inv["amount"], inv["due"], inv["late"],
              f"invoice{(' ' + inv['ref']) if inv['ref'] else ''} due {long_date(inv['due'])}; late by {inv['late']} days ({inv['late_src']})")
    for p in m["plans"]:
        for d in p["dates"]:
            if d <= end:
                r_add("Payment plans", p["client"], "payment plan", p["amount"], d, 0, "payment plan payment")
    for o in m["one_offs"]:
        if o["direction"] == "in" and o["date"] <= end:
            r_add(o["label"], "one-off", "certain" if o["certain"] else "one-off", o["amount"], o["date"], 0, "founder said")
    if m["pattern"]:
        k = 0
        while True:
            ws = start + timedelta(days=7 * k)
            if ws > end:
                break
            mid = ws + timedelta(days=3)
            if ws >= m["pattern_from"] - timedelta(days=6) and mid.month not in m.get("months_off", []):
                f = m.get("month_factor", {}).get(mid.month, 1.0)
                if f:
                    r_add("Other receipts (your usual pattern)", "pattern", "usual pattern", m["pattern"] * f, ws, 0,
                          m["pattern_src"] + (f"; {round(f * 100)}% in {MONTH_NAMES[mid.month - 1].title()}, as you said" if f != 1.0 else ""))
            k += 1

    def p_add(line, cat, amount, when, source, deadline=None, need=False):
        pay.append({"line": line, "category": cat, "amount": round(amount or 0.0, 2), "date": when, "deadline": deadline,
                    "source": source, "need": need})
    for o in m["out"]:
        for d in occurrences(o["next"], o["cadence"], end, o["until"]):
            if d >= start:
                p_add(o["label"], o["category"], o["amount"], d, o["source"])
    k = 0
    while start + timedelta(days=7 * k) <= end:
        ws = start + timedelta(days=7 * k)
        if m["everyday"]:
            p_add("Everyday spending", "everyday", m["everyday"], ws, m["everyday_src"])
        if m["owner_irreg"]:
            p_add("Owner pay (irregular)", "owner pay", m["owner_irreg"], ws, "your average")
        k += 1
    for o in m["one_offs"]:
        if o["direction"] == "out" and start <= o["date"] <= end:
            p_add(o["label"], "linked one-off" if o.get("linked") else "one-off", o["amount"], o["date"],
                  "founder said" + (f"; for {o['linked']}, so never held" if o.get("linked") else ""))
    lc = m.get("linked_cost")
    if lc:
        line = f"{lc['label']} (booked work)"
        if lc["one_off"] and start <= lc["start"] <= end:
            p_add(line, "linked cost", lc["one_off"], lc["start"], "one-off cost on the start date; tied to booked work")
        if lc["monthly"]:
            fp = first_payday(lc["start"], lc["payday"])
            k = 0
            while True:
                d = add_months(fp, k)
                if d > end or (lc.get("end") and d > lc["end"]) or k > 12:
                    break
                if d >= start:
                    p_add(line, "linked cost", lc["monthly"], d, f"monthly cost on the {lc['payday']} ({lc['payday_src']}); "
                          "tied to booked work")
                k += 1
    for k in m.get("known", []):
        when = k["dates"] or occurrences(k["next"], k["cadence"], end, k["until"])
        for d in when:
            if start <= d <= end:
                p_add(k["label"], "known", k["amount"], d, k["source"] + ("" if k["amount"] is not None else "; amount [NEED]"),
                      d, k["amount"] is None)
                taxes.append({"label": k["label"], "kind": k["kind"], "date": iso(d), "amount": k["amount"],
                              "source": k["source"], "who": k["who"]})
            elif d > end and not k["dates"]:
                break
    commit = []
    c = m.get("commit")
    if c:
        line = f"Commitment: {c['label']}"
        if c["one_off"]:
            commit.append({"line": line, "category": "commitment", "amount": round(c["one_off"], 2), "date": c["start"],
                           "source": "one-off cost on the start date", "k": None})
        if c["monthly"]:
            fp = first_payday(c["start"], c["payday"])
            k = 0
            while True:
                d = add_months(fp, k)
                if d > end + timedelta(days=7 * 13) or k > 12 or (c.get("end") and d > c["end"]):
                    break
                commit.append({"line": line, "category": "commitment", "amount": round(c["monthly"], 2), "date": d,
                               "source": f"monthly cost on the {c['payday']} ({c['payday_src']})", "k": k})
                k += 1
    return rec, pay, taxes, commit


def _ym(d):
    return d.year, d.month


# ------------------------------------------------------------------ the roll
def weeks_of(m):
    return [(k + 1, m["start"] + timedelta(days=7 * k), m["start"] + timedelta(days=7 * k + 6)) for k in range(HORIZON_WEEKS)]


def line_order(rec, pay):
    rin = []
    for kind in ("regular client", "invoice", "payment plan", "one-off", "certain", "usual pattern"):
        for e in rec:
            if e["kind"] == kind and e["line"] not in rin:
                rin.append(e["line"])
    cat_rank = {"payroll": 0, "owner pay": 1, "pension": 2, "known": 3, "rent": 4, "loan": 5, "loan or lease": 5, "credit card": 6,
                "insurance": 7, "software": 8, "phone and utilities": 9, "accountant": 10, "other regular": 11,
                "bank charges": 12, "everyday": 13, "own transfer": 14, "tax set-aside": 14, "one-off": 15,
                "linked one-off": 15, "linked cost": 3}
    totals = {}
    for e in pay:
        totals[e["line"]] = totals.get(e["line"], 0) + e["amount"]
    cats = {}
    for e in pay:
        cats.setdefault(e["line"], e["category"])
    rout = sorted(totals, key=lambda ln: (cat_rank.get(cats[ln], 11), -totals[ln]))
    return rin, rout


def roll(m, rec, pay, commit, case):
    wk = weeks_of(m)
    rin, rout = line_order(rec, pay)
    a_key, d_key = ("amount", "date") if case == "base" else ("s_amount", "s_date")
    lines_in = {ln: [0.0] * HORIZON_WEEKS for ln in rin}
    lines_out = {ln: [0.0] * HORIZON_WEEKS for ln in rout}
    for e in rec:
        for i, (_n, ws, we) in enumerate(wk):
            if ws <= e[d_key] <= we:
                lines_in[e["line"]][i] += e[a_key]
    for e in pay:
        for i, (_n, ws, we) in enumerate(wk):
            if ws <= e["date"] <= we:
                lines_out[e["line"]][i] += e["amount"]
    com = [0.0] * HORIZON_WEEKS
    for e in commit:
        for i, (_n, ws, we) in enumerate(wk):
            if ws <= e["date"] <= we:
                com[i] += e["amount"]
    weeks, bal, cbal = [], m["opening"], m["opening"]
    for i, (n, ws, we) in enumerate(wk):
        tin = round(sum(v[i] for v in lines_in.values()), 2)
        tout = round(sum(v[i] for v in lines_out.values()), 2)
        opening = bal
        closing = round(opening + tin - tout, 2)
        # the arithmetic assertion, checked against the events directly
        ev_in = sum(e[a_key] for e in rec if ws <= e[d_key] <= we)
        ev_out = sum(e["amount"] for e in pay if ws <= e["date"] <= we)
        if abs((opening + ev_in - ev_out) - closing) > 0.02:
            raise AssertionError(f"week {n}: closing {closing} is not opening {opening} + in {ev_in} less out {ev_out}")
        cbal = round(cbal + tin - tout - com[i], 2)
        weeks.append({"week": n, "start": iso(ws), "end": iso(we), "opening": round(opening, 2), "in": tin, "out": tout,
                      "closing": closing, "below": closing < m["floor"], "commit": round(com[i], 2),
                      "closing_with_commit": cbal, "below_with_commit": cbal < m["floor"]})
        bal = closing
    return {"weeks": weeks, "lines_in": {k: [round(x, 2) for x in v] for k, v in lines_in.items()},
            "lines_out": {k: [round(x, 2) for x in v] for k, v in lines_out.items()}}


def stats(case, floor, key="closing", below="below"):
    ws = case["weeks"]
    low = min(ws, key=lambda w: w[key])
    first = next((w["week"] for w in ws if w[below]), None)
    return {"lowest": low[key], "lowest_week": low["week"], "lowest_start": low["start"], "first_breach": first,
            "first_breach_start": next((w["start"] for w in ws if w["week"] == first), None),
            "runway": (first - 1) if first else HORIZON_WEEKS, "weeks_below": sum(1 for w in ws if w[below]),
            "shortfall": round(floor - low[key], 2) if low[key] < floor else 0.0, "end": ws[-1][key]}


def commitment_test(m, rec, pay, commit):
    c = m.get("commit")
    if not c:
        return None
    s = roll(m, rec, pay, commit, "stress")
    st = stats(s, m["floor"], "closing_with_commit", "below_with_commit")
    need_cushion = round(c["monthly"] * RUNWAY_MIN_WEEKS * 12 / 52, 2)
    end_margin = round(s["weeks"][-1]["closing_with_commit"] - m["floor"], 2)
    res = {"label": c["label"], "start": iso(c["start"]), "monthly": c["monthly"], "one_off": c["one_off"],
           "stress_lowest": st["lowest"], "stress_lowest_week": st["lowest_week"], "first_breach": st["first_breach"],
           "end_margin": end_margin, "cushion_rule": need_cushion}
    if st["first_breach"] is None and end_margin >= need_cushion:
        res.update(verdict="yes", later_start=None)
        return res
    if stats(s, m["floor"])["first_breach"] is not None:
        res.update(verdict="not yet", later_start=None, already=True)
        return res
    for k in range(1, HORIZON_WEEKS):
        ns = c["start"] + timedelta(days=7 * k)
        if ns > m["end"]:
            break
        mc = dict(m, commit=dict(c, start=ns))
        _r, _p, _t, cm = schedule(mc)
        s2 = roll(mc, rec, pay, cm, "stress")
        st2 = stats(s2, m["floor"], "closing_with_commit", "below_with_commit")
        if st2["first_breach"] is None and s2["weeks"][-1]["closing_with_commit"] - m["floor"] >= need_cushion:
            res.update(verdict="later", later_start=iso(ns))
            return res
    res.update(verdict="not yet", later_start=None)
    return res


# ------------------------------------------------------------------ results, actions, reading
def compute(m):
    rec, pay, taxes, commit = schedule(m)
    base = roll(m, rec, pay, commit, "base")
    stress = roll(m, rec, pay, commit, "stress")
    bs, ss = stats(base, m["floor"]), stats(stress, m["floor"])
    res = {"base": base, "stress": stress, "base_stats": bs, "stress_stats": ss, "receipts": rec, "payments": pay,
           "taxes": taxes, "commit_events": commit, "commitment": commitment_test(m, rec, pay, commit),
           "floor": m["floor"], "floor_basis": m["floor_basis"], "monthly_out": m["monthly_out"],
           "monthly_parts": m["monthly_parts"], "opening_below_floor": m["opening"] < m["floor"], "_cur": m["currency"]}
    res["urgent"] = bool(bs["first_breach"] and bs["first_breach"] <= PROXIMITY_WEEKS) or res["opening_below_floor"]
    has_payroll = any(o["category"] == "payroll" for o in m["out"]) or any(k["kind"] == "payroll tax" for k in m.get("known", []))
    res["weekly_cadence"] = res["urgent"] or ss["runway"] < RUNWAY_MIN_WEEKS or has_payroll
    res["actions"] = actions(m, res)
    return res


def poss(name):
    name = str(name).strip()
    return name + ("'" if name.endswith("s") else "'s")


_YEARS = {"v": False}  # True when the 13 weeks cross into a new year: week labels then carry the year


def wk_label(n, start_iso):
    d = from_iso(start_iso) if isinstance(start_iso, str) else start_iso
    s = short_date(d)
    if _YEARS["v"]:
        s = f"{s}, {d.year}" if re.match(r"[A-Za-z]", s) else f"{s} {d.year}"
    if _YEARS.get("date_first"):
        return f"the week from {s} (week {n})"
    return f"week {n} (from {s})"


def biggest_out_in_week(res, n, case="stress"):
    """The largest payment leaving in week n; pay lines are added together as payroll."""
    i = n - 1
    cats = {}
    for e in res["payments"]:
        cats.setdefault(e["line"], e["category"])
    groups = {}
    for ln, vals in res[case]["lines_out"].items():
        if vals[i] <= 0:
            continue
        key = "payroll" if cats.get(ln) == "payroll" else ln
        groups[key] = groups.get(key, 0) + vals[i]
    if not groups:
        return None
    k = max(groups, key=groups.get)
    return (k, round(groups[k], 2))


def actions(m, res):
    cur, today = m["currency"], m["today"]
    acts = []
    bs, ss = res["base_stats"], res["stress_stats"]
    breach = ss["first_breach"] or bs["first_breach"]
    pay = res["payments"]

    def when(d):
        return max(today, d)
    if breach or res["opening_below_floor"]:
        bw = breach or 1
        bw_end = m["start"] + timedelta(days=7 * bw - 1)
        cands = [e for e in res["receipts"] if e["kind"] in ("invoice", "regular client") and e["due"] <= bw_end + timedelta(days=21)
                 and (e["due"] < today or e["late"])]
        if cands:
            e = max(cands, key=lambda e: e["amount"])
            overdue = e["due"] < today
            wk_of = lambda d: min(HORIZON_WEEKS + 1, (d - m["start"]).days // 7 + 1)  # noqa: E731
            if e["late"]:
                lands = wk_of(e["date"])
                tail = (f": at their usual {e['late']} days late it lands in "
                        + (f"week {lands}" if lands <= HORIZON_WEEKS else "after week 13")
                        + (f", after the balance crosses the line in week {bw}." if lands > bw else "."))
            else:
                lands = wk_of(e["s_date"])
                tail = (f". In a bad quarter it lands {m['delay']} days late, in "
                        + (f"week {lands}." if lands <= HORIZON_WEEKS else "the week after this forecast ends."))
            acts.append({"date": iso(today if overdue else when(e["due"] - timedelta(days=7))), "line": e["line"],
                         "text": ((f"Chase {e['from']} for the {money(e['amount'], cur)} that was due on {long_date(e['due'])}"
                                   if overdue else
                                   f"Ask {e['from']} to pay the {money(e['amount'], cur)} due on {long_date(e['due'])} on time")
                                  + tail)})
        elif m["pattern"]:
            acts.append({"date": iso(today), "line": "Other receipts (your usual pattern)",
                         "text": f"Collect the money for work in {('weeks ' + str(max(1, bw - 2)) + ' to ' + str(bw)) if bw > 1 else 'week 1'} on time (invoices out, deposits "
                                 f"and balances in): your usual receipts of {money(m['pattern'], cur)} a week are what carry "
                                 f"week {bw}, and in a bad quarter they arrive {m['delay']} days later."})
        low_w = max(bs["lowest_week"] if bs["first_breach"] else ss["lowest_week"], bw)
        win_end = m["start"] + timedelta(days=7 * min(low_w, bw + 3) - 1)
        wk_pay = [e for e in pay if m["start"] + timedelta(days=7 * (bw - 2)) <= e["date"] <= win_end and e["amount"] > 0
                  and e["date"] >= today]
        known_kind = {k["label"]: k["kind"] for k in m.get("known", [])}
        # what can be held, in order: money moved to the founder's own accounts, then owner pay, then spending that
        # is not tied to booked work. Payroll taxes, pay, and payments tied to a booked job or invoice are never held.
        tiers = [[e for e in wk_pay if e["category"] in ("own transfer",)],
                 [e for e in wk_pay if e["category"] == "owner pay"],
                 [e for e in wk_pay if e["category"] in ("one-off", "other regular", "software", "everyday")]]
        tax = [e for e in wk_pay if e["category"] == "known" and known_kind.get(e["line"]) != "payroll tax" and not e["need"]]
        first_tier = next((t for t in tiers if t), [])
        holdable = sorted((e for t in tiers for e in t), key=lambda e: -e["amount"])
        if tiers[0] or tiers[1]:
            e = max(first_tier, key=lambda e: e["amount"])
            acts.append(hold_move(e, m, today))
        elif tax and (not tiers[2] or max(tax, key=lambda e: e["amount"])["amount"] >= max(tiers[2], key=lambda e: e["amount"])["amount"]):
            e = max(tax, key=lambda e: e["amount"])
            dl = e.get("deadline") or e["date"]
            who = next((k["who"] for k in m.get("known", []) if k["label"] == e["line"]), "your accountant")
            who = who if re.search(r"accountant|bookkeeper|payroll|adviser|advisor|lawyer|cpa", who, re.I) else "your accountant"
            acts.append({"date": iso(when(dl - timedelta(days=14))), "line": e["line"],
                         "text": f"Talk to {who} by {long_date(when(dl - timedelta(days=14)))} about the "
                                 f"{money(e['amount'], cur)} {e['line']} due {long_date(dl)}. It lands while the balance "
                                 "is under the line."})
        elif tiers[2]:
            acts.append(hold_move(max(tiers[2], key=lambda e: e["amount"]), m, today))
    else:
        tx = [e for e in pay if e["category"] == "known"]
        if tx:
            e = min(tx, key=lambda e: e["date"])
            who = next((k["who"] for k in m.get("known", []) if k["label"] == e["line"]), "")
            who = who if re.search(r"accountant|bookkeeper|payroll|adviser|advisor|lawyer|cpa", who, re.I) else "your accountant"
            if e["need"]:
                acts.append({"date": iso(when(e["date"] - timedelta(days=21))), "line": e["line"],
                             "text": f"Get the {e['line']} figure due {long_date(e.get('deadline') or e['date'])} from {who} by "
                                     f"{long_date(when(e['date'] - timedelta(days=21)))}; it is placed at zero until then."})
            else:
                acts.append({"date": iso(when(e["date"] - timedelta(days=7))), "line": e["line"],
                             "text": f"Move {money(e['amount'], cur)} for the {e['line']} due {long_date(e.get('deadline') or e['date'])} "
                                     f"into a separate account by {long_date(when(e['date'] - timedelta(days=7)))}, so it is never "
                                     "counted as spare cash."})
        big = max((e for e in res["receipts"] if e["kind"] != "usual pattern"), key=lambda e: e["amount"], default=None)
        if big:
            acts.append({"date": iso(when(big["due"] - timedelta(days=5))), "line": big["line"],
                         "text": f"Confirm {poss(big['from'])} {money(big['amount'], cur)} due {long_date(big['due'])}: it is the "
                                 "largest single receipt in the 13 weeks."})
    c = res.get("commitment")
    if c and c["verdict"] != "yes":
        line = f"Commitment: {c['label']}"
        if c["verdict"] == "later":
            acts.append({"date": iso(today), "line": line,
                         "text": f"Move the {c['label']} start from {long_date(c['start'])} to {long_date(c['later_start'])}: "
                                 "the earliest start that keeps a bad quarter above the safety line."})
        else:
            acts.append({"date": iso(today), "line": line,
                         "text": f"Hold off on the {c['label']}: " + ("a bad quarter breaches the line even without it. "
                                 if c.get("already") else "no start inside these 13 weeks keeps a bad quarter above the "
                                 "safety line. ") + "Re-test it on the next run."})
    nxt = today + timedelta(days=7 if res["weekly_cadence"] else 28)
    rerun = {"date": iso(nxt), "line": "Opening balance",
             "text": f"Re-run this forecast with a fresh export on {long_date(nxt)} "
                     + ("(weekly while a bad quarter crosses the line within 8 weeks, or while you run payroll)."
                        if res["weekly_cadence"] else "(monthly, and before any new commitment).")}
    res["next_run"] = rerun
    # the re-run date is printed on its own line; the three moves are things to do, each with a figure
    for x in fill_moves(m, res, holdable if (breach or res["opening_below_floor"]) else []):
        acts.append(x)
    # the founder's own priority wins: moves on the lines they named come first and are kept
    first = [str(x).lower() for x in (m.get("moves_first") or []) if str(x).strip()]
    pin = lambda x: any(f in x["line"].lower() or x["line"].lower() in f for f in first)  # noqa: E731
    acts = [x for x in acts if pin(x)] + [x for x in acts if not pin(x)]
    seen, out = set(), []
    for x in acts:
        fig = re.search(r"\d[\d,. ]*\d|\d", x["text"])
        key = (x["line"], fig.group(0) if fig and x["line"].startswith("Invoices:") else "")
        if key not in seen:
            out.append(x)
            seen.add(key)
    out = out[:3]
    if len(out) < 3:
        out.append(rerun)
    return sorted(out, key=lambda x: (not pin(x), x["date"]))


def fill_moves(m, res, holdable):
    """Further moves, in order of use, so there are three real ones before a re-run is ever counted as a move:
    getting a receipt that lands after the low paid on time, a second payment that can wait, money set aside for a
    payment the founder listed, the largest receipt to confirm, and only then the gap itself."""
    cur, today = m["currency"], m["today"]
    out = []
    bs, ss = res["base_stats"], res["stress_stats"]
    case, st = ("normal", bs) if bs["first_breach"] else ("bad", ss)
    lw_start = from_iso(st["lowest_start"])
    lw_end = lw_start + timedelta(days=6)
    if st["shortfall"] > 0:
        key_d = "date" if case == "normal" else "s_date"
        later = [e for e in res["receipts"] if e["kind"] in ("invoice", "regular client") and e["late"]
                 and e[key_d] > lw_end and e["due"] <= lw_end]
        if later:
            r = max(later, key=lambda e: e["amount"])
            by = max(today, r["due"] - timedelta(days=7))
            out.append({"date": iso(by), "line": r["line"],
                        "text": f"Ask {r['from']} to pay the {money(r['amount'], cur)} due {long_date(r['due'])} on time: at "
                                f"their usual {r['late']} days late it lands after "
                                f"the {case}-quarter low of {money(st['lowest'], cur)} in "
                                f"{wk_label(st['lowest_week'], st['lowest_start'])}, "
                                f"{money(st['shortfall'], cur)} under the line."})
    for e in [e for e in holdable if e["category"] in ("own transfer", "owner pay", "one-off")][:3]:
        out.append(hold_move(e, m, today))
    kinds = {k["label"]: k["kind"] for k in m.get("known", [])}
    tx = sorted((e for e in res["payments"] if e["category"] == "known" and not e["need"] and e["date"] >= today
                 and kinds.get(e["line"]) != "payroll tax"), key=lambda e: e["date"])
    if tx:
        e = tx[0]
        by = max(today, e["date"] - timedelta(days=7))
        out.append({"date": iso(by), "line": e["line"],
                    "text": f"Move {money(e['amount'], cur)} for the {e['line']} due {long_date(e.get('deadline') or e['date'])} "
                            f"into a separate account by {long_date(by)}, so it is never counted as spare cash."})
    big = max((e for e in res["receipts"] if e["kind"] != "usual pattern"), key=lambda e: e["amount"], default=None)
    if big:
        by = max(today, big["due"] - timedelta(days=5))
        out.append({"date": iso(by), "line": big["line"],
                    "text": f"Confirm {poss(big['from'])} {money(big['amount'], cur)} due {long_date(big['due'])}: it is the "
                            "largest single receipt in the 13 weeks."})
    if st["shortfall"] > 0:
        by = max(today, lw_start - timedelta(days=7))
        out.append({"date": iso(by), "line": "Safety line",
                    "text": f"Find {money(st['shortfall'], cur)} by {long_date(by)}: that is how far the {case}-quarter low of "
                            f"{money(st['lowest'], cur)} in {wk_label(st['lowest_week'], st['lowest_start'])} sits under the "
                            f"{money(m['floor'], cur)} line, as new money in or a cost cut."})
    return out


def decides(m, res):
    """The receipt that decides the quarter, in plain arithmetic on the forecast's own weeks. First choice: a late
    payer whose lateness moves the result (what the quarter looks like if it paid on its due date). Otherwise: the
    last date a large receipt can land before a normal quarter falls below the line, or, when it is already below,
    before the account is overdrawn. Only said when it changes the answer; nothing is estimated."""
    cur = m["currency"]
    ws = res["base"]["weeks"]
    floor = m["floor"]
    base_first = next((w["week"] for w in ws if w["closing"] < floor), None)
    base_zero = next((w["week"] for w in ws if w["closing"] < 0), None)
    wk = lambda d: (d - m["start"]).days // 7 + 1  # noqa: E731
    rec = [e for e in res["receipts"] if e["kind"] != "usual pattern" and m["start"] <= e["date"] <= m["end"]]
    if not rec:
        return None

    def what(e):
        return (f"{poss(e['from'])} {money(e['amount'], cur)}" if e["from"] not in ("one-off", "pattern")
                else f"{e['line']} ({money(e['amount'], cur)})")
    # 1. a late payer whose lateness changes the week
    for e in sorted((e for e in rec if e["late"] and e["due"] >= m["start"]), key=lambda e: -e["amount"]):
        kd, k = wk(e["due"]), wk(e["date"])
        if kd >= k:
            continue
        alt = [w["closing"] + (e["amount"] if kd <= w["week"] < k else 0) for w in ws]
        alt_first = next((w["week"] for w, c in zip(ws, alt) if c < floor), None)
        alt_zero = next((w["week"] for w, c in zip(ws, alt) if c < 0), None)
        head = (f"{what(e)} decides it: at their usual {e['late']} days late it lands in week {k}; paid on its due date, "
                f"{long_date(e['due'])}, ")
        if base_first and alt_first != base_first:
            return head + ("a normal quarter would stay above the line for all 13 weeks." if alt_first is None else
                           f"a normal quarter would first fall below the line in week {alt_first} instead of week {base_first}.")
        if base_zero and alt_zero != base_zero:
            return head + ("the account would not go overdrawn in these 13 weeks." if alt_zero is None else
                           f"the account would first go overdrawn in week {alt_zero} instead of week {base_zero}.")
    # 2. the last date a large receipt can land
    top = max(e["amount"] for e in rec)
    cands = sorted((e for e in rec if e["amount"] >= 0.75 * top), key=lambda e: e["date"]) + sorted(
        (e for e in rec if e["amount"] < 0.75 * top), key=lambda e: -e["amount"])
    for e in cands[:6]:
        k = wk(e["date"])
        without = [w["closing"] - (e["amount"] if w["week"] >= k else 0) for w in ws]
        tip = next((w for w, c in zip(ws, without) if c < floor), None)
        zero = next((w for w, c in zip(ws, without) if c < 0), None)
        if tip is None:
            return (f"{what(e)} is the largest single receipt, due to land in week {k}. Even if it does not land in these 13 "
                    f"weeks, a normal quarter stays above the line (lowest {money(min(without), cur)}).")
        if tip["week"] >= k and (base_first is None or tip["week"] < base_first):
            txt = (f"{what(e)} decides the quarter: if it has not landed by {long_date(tip['end'])}, the end of "
                   f"{wk_label(tip['week'], tip['start'])}, a normal quarter falls below the line that week")
            if zero is not None and (base_zero is None or zero["week"] < base_zero):
                txt += f", and the account is overdrawn in {wk_label(zero['week'], zero['start'])} if it has still not landed by then"
            return txt + "."
        if zero is not None and zero["week"] >= k and (base_zero is None or zero["week"] < base_zero):
            return (f"{what(e)} decides how far it falls: if it has not landed by {long_date(zero['end'])}, the end of "
                    f"{wk_label(zero['week'], zero['start'])}, the account is overdrawn that week.")
    return None


def hold_move(e, m, today):
    cur = m["currency"]
    d = max(today, e["date"] - timedelta(days=3))
    return {"date": iso(d), "line": e["line"],
            "text": f"Hold the {e['line']} of {money(e['amount'], cur)} due {long_date(e['date'])} until the week the "
                    f"balance is back above {money(m['floor'], cur)}, or cut it."}


def headline(m, res):
    cur = m["currency"]
    bs, ss = res["base_stats"], res["stress_stats"]
    fl = money(m["floor"], cur)
    if res["opening_below_floor"]:
        ws = res["base"]["weeks"]
        back = next((w for w in ws if not w["below"]), None)
        again = next((w for w in ws if back and w["week"] > back["week"] and w["below"]), None)
        if not back:
            tail = "in a normal quarter it does not climb back above the line in these 13 weeks"
        elif again:
            tail = (f"in a normal quarter it climbs back above the line in {wk_label(back['week'], back['start'])} and "
                    f"falls below it again in {wk_label(again['week'], again['start'])}")
        else:
            tail = f"in a normal quarter it climbs back above the line in {wk_label(back['week'], back['start'])} and stays there"
        return (f"The balance of {money(m['opening'], cur)} on {long_date(m['opening_date'])} is already below the {fl} "
                f"safety line ({m['floor_basis']}), and {tail}. At its lowest in a normal quarter the account "
                f"{bal_phrase(bs['lowest'], cur)}, in {wk_label(bs['lowest_week'], bs['lowest_start'])}.")

    def why(n, case):
        b = biggest_out_in_week(res, n, case)
        if not b:
            return ""
        w = res[case]["weeks"][n - 1]
        if b[1] < w["opening"] + w["in"] - m["floor"]:
            # no single payment takes it through the line: say what the week does as a whole
            big = "payroll of " + money(b[1], cur) if b[0] == "payroll" else f"{b[0]} ({money(b[1], cur)})"
            if n == 1 and m["opening"] >= m["floor"]:
                return (f", because the balance starts only {money(m['opening'] - m['floor'], cur)} above it and week 1 has "
                        f"{money(w['out'], cur)} going out against {money(w['in'], cur)} coming in")
            return f", when {money(w['out'], cur)} goes out against {money(w['in'], cur)} coming in, {big} the largest part"
        return f", when {'payroll of ' + money(b[1], cur) if b[0] == 'payroll' else b[0] + ' (' + money(b[1], cur) + ')'} leaves"
    if bs["first_breach"] and ss["first_breach"]:
        low = (f" At its lowest in a normal quarter the account {bal_phrase(bs['lowest'], cur)}, in "
               f"{wk_label(bs['lowest_week'], bs['lowest_start'])}." if bs["lowest_week"] != bs["first_breach"] else "")
        if bs["first_breach"] == ss["first_breach"]:
            return (f"Your cash falls below its {fl} safety line in {wk_label(bs['first_breach'], bs['first_breach_start'])}, "
                    f"in a normal quarter and a bad one{why(bs['first_breach'], 'base')}.{low}")
        return (f"Your cash falls below its {fl} safety line in {wk_label(bs['first_breach'], bs['first_breach_start'])} "
                f"in a normal quarter{why(bs['first_breach'], 'base')}, and in "
                f"{wk_label(ss['first_breach'], ss['first_breach_start'])} in a bad one.{low}")
    if ss["first_breach"]:
        return (f"In a normal quarter your cash stays above its {fl} safety line, lowest at {money(bs['lowest'], cur)} in "
                f"{wk_label(bs['lowest_week'], bs['lowest_start'])}. In a bad quarter it drops below it in "
                f"{wk_label(ss['first_breach'], ss['first_breach_start'])}{why(ss['first_breach'], 'stress')}.")
    return (f"Your cash stays above its {fl} safety line for all 13 weeks, in a normal quarter and a bad one. "
            f"The lowest point in a bad quarter is {money(ss['lowest'], cur)} in "
            f"{wk_label(ss['lowest_week'], ss['lowest_start'])}.")


def bal_phrase(v, cur=None):
    return f"is {money(-v, cur)} overdrawn" if v < 0 else f"holds {money(v, cur)}"


def overdrawn_line(res):
    out = []
    for case, word in (("base", "a normal quarter"), ("stress", "a bad quarter")):
        neg = [w for w in res[case]["weeks"] if w["closing"] < 0]
        if neg:
            low = min(neg, key=lambda w: w["closing"])
            out.append(f"In {word} the account is overdrawn from {wk_label(neg[0]['week'], neg[0]['start'])}, "
                       f"by as much as {money(-low['closing'], res['_cur'])} in week {low['week']}.")
    return " ".join(out) or None


def runway_line(res):
    r = res["stress_stats"]["runway"]
    if res["opening_below_floor"]:
        return "Bad-quarter runway: none. The balance starts below the safety line, against an 8-week minimum."
    if res["stress_stats"]["first_breach"] is None:
        return f"Bad-quarter runway: no breach in 13 weeks (the method asks for at least {RUNWAY_MIN_WEEKS})."
    word = "meets" if r >= RUNWAY_MIN_WEEKS else "is below"
    return (f"Bad-quarter runway: {r} week{'s' if r != 1 else ''} before the balance crosses the safety line, which "
            f"{word} the {RUNWAY_MIN_WEEKS}-week minimum.")


def commitment_line(m, res):
    c = res.get("commitment")
    if not c:
        return None
    cur = m["currency"]
    cost = (f"{money(c['monthly'], cur)} a month" if c["monthly"] else "") + (
        f"{' plus ' if c['monthly'] else ''}{money(c['one_off'], cur)} up front" if c["one_off"] else "")
    if c["verdict"] == "yes":
        return (f"{c['label']} from {long_date(c['start'])} ({cost}): yes. With it loaded, a bad quarter stays above "
                f"the safety line for all 13 weeks (lowest {money(c['stress_lowest'], cur)} in week {c['stress_lowest_week']}) "
                f"and ends {money(c['end_margin'], cur)} above it, more than {RUNWAY_MIN_WEEKS} weeks of the cost.")
    why = (f"A bad quarter crosses the safety line in week {c['first_breach']}." if c["first_breach"] else
           f"A bad quarter stays above the line, but ends week 13 only {money(c['end_margin'], cur)} above it, less than "
           f"{RUNWAY_MIN_WEEKS} weeks of the cost ({money(c['cushion_rule'], cur)}).")
    if c["verdict"] == "later":
        return (f"{c['label']} from {long_date(c['start'])} ({cost}): not yet. {why} Starting on "
                f"{long_date(c['later_start'])} keeps it above the line through week 13, with at least "
                f"{RUNWAY_MIN_WEEKS} weeks of the cost ({money(c['cushion_rule'], cur)}) in hand at the end. Test it "
                f"again before you commit: this forecast cannot see past {long_date(m['end'])}.")
    if c.get("already"):
        return (f"{c['label']} from {long_date(c['start'])} ({cost}): not yet. A bad quarter already crosses the safety "
                "line without it, so no start date inside these 13 weeks can pass. Fix the breach first, then test it again.")
    return (f"{c['label']} from {long_date(c['start'])} ({cost}): not yet. {why} No start inside these 13 weeks clears it.")


def known_lines_text(m, res):
    """The payments the founder listed, one line each, with every date inside the window and who gave the figure."""
    cur = m["currency"]
    out, groups = [], {}
    for t in res["taxes"]:
        groups.setdefault((t["label"], t["amount"], t["source"]), []).append(t["date"])
    for (label, a, src), dates in groups.items():
        fig = money(a, cur) if a is not None else "[NEED: amount from you or your accountant]"
        ds = [long_date(d) for d in sorted(dates)]
        when = ds[0] if len(ds) == 1 else ", ".join(ds[:-1]) + " and " + ds[-1]
        out.append(f"{label}: {fig}{' each' if len(ds) > 1 and a is not None else ''} on {when}; {src}.")
    for k in m.get("known", []):
        if not any(t["label"] == k["label"] for t in res["taxes"]):
            when = k["dates"][0] if k["dates"] else k["next"]
            out.append(f"{k['label']}: {long_date(when)}, outside these 13 weeks; {k['source']}.")
    return out


def pattern_line(m):
    """How the usual pattern of other receipts was built, in one sentence, when it is in the forecast."""
    if not m.get("pattern") and not m.get("pattern_removed"):
        return None
    cur = m["currency"]
    bits = [f"Other receipts: your usual pattern of {money(m['pattern'], cur)} a week"
            + (f", the average of {m['pattern_months']}" if m.get("pattern_months") else "")]
    if m.get("pattern_removed"):
        bits.append(f"without {' and '.join(m['pattern_removed'])}, whose money in comes from the invoices listed")
    if m.get("months_off"):
        bits.append("none in " + " and ".join(MONTH_NAMES[k - 1].title() for k in m["months_off"]) + ", as you said")
    if m.get("month_factor"):
        bits.append("; ".join(f"{round(v * 100)}% of it in {MONTH_NAMES[k - 1].title()}" for k, v in sorted(m["month_factor"].items())))
    if m["invoices"]:
        bits.append(f"starting {long_date(m['pattern_from'])}, when new work at your usual terms is paid")
    return ", ".join(bits) + "."


def horizon_lines(m):
    """Say plainly when the month the founder asked about runs past what the 13 weeks can see."""
    out = []
    for name in m.get("asked_about") or []:
        mo = months_of([name])
        if not mo:
            continue
        k = mo[0]
        y = m["today"].year if k >= m["today"].month else m["today"].year + 1
        first_day, last_day = date(y, k, 1), month_end(y, k)
        title = MONTH_NAMES[k - 1].title()
        if last_day <= m["end"]:
            continue
        if first_day > m["end"]:
            out.append(f"You asked about {title}. This forecast runs to {long_date(m['end'])}, so {title} is past what it "
                       f"can see. Re-run it in the first week of {MONTH_NAMES[(k - 2) % 12].title()} to see {title}.")
        else:
            out.append(f"You asked about {title}. This forecast runs to {long_date(m['end'])}, so the rest of {title} is "
                       f"past what it can see. Re-run it in the first week of {title} to see the rest.")
    return out


def reading_md(m, res, needs, notes, note_text=None, variance=None):
    cur = m["currency"]
    bs, ss = res["base_stats"], res["stress_stats"]
    L = [f"# 13-week cash forecast: {m['business']}", "",
         f"Made {long_date(m['today'])}. Weeks run from {long_date(m['start'])} to {long_date(m['end'])}.", ""]
    moves = ["## Three moves", ""] + [f"{i}. **By {long_date(a['date'])}**, on the line \"{a['line']}\": {a['text']}"
                                      for i, a in enumerate(res["actions"], 1)]
    if res["next_run"] not in res["actions"]:
        moves += ["", f"Next run: {res['next_run']['text']}"]
    moves.append("")
    if m.get("moves_on_top"):
        L += moves
    L += [f"**{headline(m, res)}**", ""]
    cl = commitment_line(m, res)
    if cl:
        L += [f"**The decision.** {cl}", ""]
    L += [runway_line(res) + (" " + overdrawn_line(res) if overdrawn_line(res) else ""), ""]
    dc = decides(m, res)
    if dc:
        L += [dc, ""]
    if note_text:
        L += ["## What this means", "", note_text.strip(), ""]
    L += ["## The figures", "", "| | Normal quarter | Bad quarter |", "|---|---|---|",
          f"| Opening balance | {money(m['opening'], cur)} on {long_date(m['opening_date'])} | same |",
          f"| Safety line | {money(m['floor'], cur)} ({m['floor_basis']}) | same |",
          f"| Lowest point | {money(bs['lowest'], cur)}, {wk_label(bs['lowest_week'], bs['lowest_start'])} | "
          f"{money(ss['lowest'], cur)}, {wk_label(ss['lowest_week'], ss['lowest_start'])} |",
          f"| First week below the line | {wk_label(bs['first_breach'], bs['first_breach_start']) if bs['first_breach'] else 'no breach in 13 weeks'} | "
          f"{wk_label(ss['first_breach'], ss['first_breach_start']) if ss['first_breach'] else 'no breach in 13 weeks'} |",
          f"| Balance at week 13 | {money(bs['end'], cur)} | {money(ss['end'], cur)} |", "",
          f"Opening balance: {m['opening_source']}. A bad quarter means money coming in is {round(m['cut'] * 100)}% lower and "
          f"{m['delay']} days later, and {round(m['plan_default'] * 100)}% of payment plan payments fail; what goes out stays the same.", ""]
    pl = pattern_line(m)
    if pl:
        L += [pl, ""]
    for h in horizon_lines(m):
        L += [h, ""]
    if not m.get("moves_on_top"):
        L += moves
    tl = known_lines_text(m, res)
    if tl:
        L += ["## Payments you listed", ""] + [f"- {x}" for x in tl] + [""]
    if m.get("tax_seen"):
        seen = m["tax_seen"][:5]
        L += ["Payments in your export that looked like tax or government payments. They are not projected; only the "
              "payments you listed above are in the forecast:", ""]
        for t in seen:
            n = int(t.get("times") or 1)
            lo, hi = t.get("low", t.get("last_amount")), t.get("high", t.get("last_amount"))
            fig = (f"{money(lo, cur)} to {money(hi, cur)}" if lo is not None and hi is not None and lo != hi
                   else money(t.get("last_amount"), cur))
            L.append(f"- {t.get('label')}: {n} payment{'s' if n != 1 else ''}, {fig}, last on {long_date(t.get('last_date'))}")
        if len(m["tax_seen"]) > 5:
            L.append(f"- and {len(m['tax_seen']) - 5} more")
        L.append("")
    if variance:
        L += ["## Since last time", "", variance, ""]
    if needs:
        L += ["## Gaps", ""] + [f"- [NEED] {x}" for x in needs] + [""]
    L += ["Planning, not advice. Tax, payroll and other government payments are the figures and dates you gave "
          "(founder supplied); confirm them with your accountant or payroll provider. This forecast does no tax "
          "arithmetic. Borrowing and employment costs: speak to your accountant.", ""]
    return "\n".join(L)


# ------------------------------------------------------------------ the note gate
_MON = (r"(January|February|March|April|May|June|July|August|September|October|November|December|"
        r"Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)")
DATE_RE = re.compile(r"\b(\d{1,2})(?:st|nd|rd|th)? " + _MON + r"\b( \d{4})?", re.I)
DATE_MF_RE = re.compile(r"\b" + _MON + r"\.? (\d{1,2})(?:st|nd|rd|th)?\b(,? \d{4})?", re.I)


def _dates_in(text):
    """Every (day, month) written in either order, and the text with them blanked out."""
    found = [(int(d), mo[:3].lower()) for d, mo, _y in DATE_RE.findall(text)]
    text = DATE_RE.sub(" ", text)
    found += [(int(d), mo[:3].lower()) for mo, d, _y in DATE_MF_RE.findall(text)]
    return found, DATE_MF_RE.sub(" ", text)


def _token_re():
    """Figures as this founder writes them (1,234.56 or 1 234,56 or 1.234,56), with a currency sign or code."""
    g, d = re.escape(NUM["group"]), re.escape(NUM["decimal"])
    syms = sorted({v for v in SYMBOLS.values()} | {"$", "US$", "S$", "SG$", "A$", "AU$", "NZ$", "C$", "CA$", "HK$"},
                  key=len, reverse=True)
    sym = "|".join(re.escape(x) for x in syms)
    num = rf"\d{{1,3}}(?:{g}\d{{3}}(?!\d))+(?:{d}\d+)?|\d+(?:{d}\d+)?"
    return re.compile(rf"(-?)((?:{sym})\s?)?({num})(\s?(?:%|per cent))?(\s?(?:{sym}|[A-Z]{{3}})\b)?"), g, d


def _figures(text):
    rx, g, d = _token_re()
    out = []
    for mm in rx.finditer(text):
        raw = mm.group(3)
        try:
            v = round(float(raw.replace(NUM["group"], "").replace(NUM["decimal"], ".")), 2)
        except ValueError:
            continue
        kind = "money" if (mm.group(2) or mm.group(5)) else ("pct" if mm.group(4) else "plain")
        out.append((mm.group(0).strip(), v, kind))
    return out


def check_note(note, allowed_text, cur=None):
    import check_copy
    problems = []
    allowed_dates, allowed_text = _dates_in(allowed_text)
    note_dates, stripped = _dates_in(note)
    for d, mo in note_dates:
        if (d, mo) not in set(allowed_dates):
            problems.append(f"date {d} {mo} is not a date in the forecast")
    stripped = re.sub(r"\bweeks? (\d{1,2})\b", lambda mm: " " if 1 <= int(mm.group(1)) <= 13 else mm.group(0), stripped, flags=re.I)
    stripped = re.sub(r"\b(\d{1,2}) (weeks?|days?|months?)\b", " ", stripped, flags=re.I)
    allowed = {}
    for _t, v, kind in _figures(allowed_text):
        allowed.setdefault(kind, set()).add(v)
    every = set().union(*allowed.values()) if allowed else set()
    for tok, v, kind in _figures(stripped):
        if kind == "plain" and v in (1, 2, 3) and float(v).is_integer():
            continue
        ok = v in allowed.get(kind, set()) if kind in ("money", "pct") else v in every
        if not ok:
            problems.append(f"figure {tok} is not in the forecast")
    found, _w, _n = check_copy.check(note, currency=cur)
    problems += [x for x in found if "Merigrove" not in x]
    words = len(re.findall(r"\b\w[\w'-]*\b", note))
    if words > 120:
        problems.append(f"note is {words} words, limit 120")
    return problems


# ------------------------------------------------------------------ since last time
def variance_text(m, prev):
    """Compare what came in since the last run with what that run forecast, from the saved projection."""
    if not prev:
        return None, None
    proj = prev.get("_projection") or []
    hist = m.get("history_daily") or {}
    if not proj or not hist:
        return None, None
    last_date = prev.get("_opening_date")
    if not last_date:
        return None, None
    lo = from_iso(last_date) + timedelta(days=1)
    hi = m["opening_date"]
    weeks_in = [p for p in proj if from_iso(p["week_end"]) <= hi]
    if not weeks_in or hi <= lo:
        return None, None
    upto = from_iso(weeks_in[-1]["week_end"])
    forecast_in = sum(float(p["receipts"]) for p in weeks_in)
    actual_in = sum(v[0] for d, v in hist.items() if lo <= from_iso(d) <= upto)
    if forecast_in <= 0:
        return None, None
    diff = (actual_in - forecast_in) / forecast_in
    cur = m["currency"]
    txt = (f"From {long_date(lo)} to {long_date(upto)} the last forecast expected {money(forecast_in, cur)} in; "
           f"{money(actual_in, cur)} arrived ({'+' if diff >= 0 else ''}{round(diff * 100)}%).")
    return txt, {"from": iso(lo), "to": iso(upto), "forecast": round(forecast_in, 2), "actual": round(actual_in, 2),
                 "diff": round(diff, 4)}


def stress_suggestion(prev, var):
    checks = list((prev or {}).get("_checks") or [])
    if var:
        checks.append(var)
    if len(checks) < 3:
        return None
    med = statistics.median(float(c["diff"]) for c in checks[-6:])
    if med < -DEFAULT_STRESS_CUT:
        newcut = min(0.5, round((-med + 0.05) * 20) / 20)
        return (f"Over your last {len(checks[-6:])} runs, money in came {round(-med * 100)}% under the normal-quarter forecast "
                f"(the middle figure). Propose a bad-quarter cut of {round(newcut * 100)}% in place of "
                f"{round(DEFAULT_STRESS_CUT * 100)}%, as a learning for the founder to approve.")
    return None


# ------------------------------------------------------------------ workbook
YELLOW = "FFF2CC"


def write_xlsx(path, m, res, reading_lines):
    import openpyxl as ox  # optional: without it the build writes CSV files instead
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as COL
    wb = ox.Workbook()
    wb.properties.creator = "Made with Merigrove " + SKILL_NAME
    wb.properties.title = f"13-week cash forecast: {m['business']}"
    bold = Font(bold=True)
    head = Font(bold=True, size=14)
    fill = PatternFill("solid", fgColor=YELLOW)
    sym = SYMBOLS.get(m["currency"] or "", "")
    after = (m["currency"] in SYMBOL_AFTER) or (m.get("number_style") or {}).get("after")
    if sym and after:
        cfmt = f'#,##0" {sym}";[Red]-#,##0" {sym}";"-"'
    elif sym:
        cfmt = f'"{sym}"#,##0;[Red]-"{sym}"#,##0;"-"'
    else:
        cfmt = '#,##0;[Red]-#,##0;"-"'  # the viewer's own thousands and decimal marks apply in the spreadsheet app
    month_first = m.get("date_order") == "month-first"
    dfmt = "mmm d, yyyy" if month_first else "d mmm yyyy"
    wfmt = "mmm d" if month_first else "d mmm"
    summ = wb.active
    summ.title = "Summary"
    ws_base = wb.create_sheet("Weekly")
    ws_st = wb.create_sheet("Stress")
    asum = wb.create_sheet("Assumptions")
    known = wb.create_sheet("Known Dates")
    recs = wb.create_sheet("Receipts Schedule")

    # Assumptions
    asum["A1"] = "Assumptions (yellow cells are yours to change; every sheet follows them)"
    asum["A1"].font = head
    rows = [("Opening balance", m["opening"], cfmt, True, m["opening_source"]),
            ("Balance date", m["opening_date"], dfmt, True, "the date of the opening balance"),
            ("Week 1 starts", "=B4+1", dfmt, False, "the day after the balance date"),
            ("Safety line (the floor)", m["floor"], cfmt, True, m["floor_basis"] + "; " + "; ".join(
                f"{k} {money(v, m['currency'])}" for k, v in m["monthly_parts"])),
            ("Monthly outgoings the line is built from", m["monthly_out"], cfmt, False,
             "regular bills, pay, the payments you listed that come monthly or more often, everyday spending; not "
             "quarterly or yearly payments, transfers to your own accounts or one-offs"),
            ("Bad quarter: money in lower by", m["cut"], "0%", True, "working default 20%"),
            ("Bad quarter: money in later by (days)", m["delay"], "0", True, "working default 15 days"),
            ("Bad quarter: payment plan payments that fail", m["plan_default"], "0%", True, "working default 10% until you have history"),
            ("Bad-quarter runway the method asks for (weeks)", RUNWAY_MIN_WEEKS, "0", False, "Merigrove method"),
            ]
    c = m.get("commit")
    if c:
        rows += [("Commitment: first monthly payment", first_payday(c["start"], c["payday"]) if c["monthly"] else c["start"],
                  dfmt, True, f"{c['label']}, start {long_date(c['start'])}; pay day {c['payday']} ({c['payday_src']})"),
                 ("Commitment: monthly cost", c["monthly"], cfmt, True, f"pay plus employer costs ({c['on_costs_src']})"),
                 ("Commitment: one-off cost on the start date", c["one_off"], cfmt, True, "recruitment, equipment"),
                 ("Commitment: start date", c["start"], dfmt, True, "")]
    for i, (lab, val, fmt, yellow, note) in enumerate(rows, start=3):
        asum.cell(i, 1, lab)
        cl = asum.cell(i, 2, val)
        cl.number_format = fmt
        if yellow:
            cl.fill = fill
        asum.cell(i, 3, note)
    A = {lab: f"Assumptions!$B${i}" for i, (lab, *_r) in enumerate(rows, start=3)}
    r = 3 + len(rows) + 1
    asum.cell(r, 1, "Currency").font = bold
    asum.cell(r, 2, m["currency"] or "[NEED] not confirmed")
    r += 1
    asum.cell(r, 1, "Dates written").font = bold
    asum.cell(r, 2, {"day-first": "day first", "month-first": "month first"}.get(m.get("date_order"), "not set"))
    r += 2
    asum.cell(r, 1, "Planning, not advice. Tax, payroll and other government payments are founder supplied; "
                    "confirm them with your accountant or payroll provider.").font = bold
    asum.column_dimensions["A"].width = 48
    asum.column_dimensions["B"].width = 16
    asum.column_dimensions["C"].width = 90

    # Receipts Schedule
    hdr = ["Line", "From", "Kind", "Amount", "Due date", "Days late", "Expected date (normal quarter)",
           "Bad quarter amount", "Bad quarter date", "Basis"]
    for j, h in enumerate(hdr, 1):
        recs.cell(1, j, h).font = bold
    start_ref, cut_ref, delay_ref, plan_ref = (A["Week 1 starts"], A["Bad quarter: money in lower by"],
                                               A["Bad quarter: money in later by (days)"],
                                               A["Bad quarter: payment plan payments that fail"])
    for i, e in enumerate(res["receipts"], start=2):
        recs.cell(i, 1, e["line"])
        recs.cell(i, 2, str(e["from"])[:60])
        recs.cell(i, 3, e["kind"])
        x = recs.cell(i, 4, e["amount"]); x.number_format = cfmt; x.fill = fill
        x = recs.cell(i, 5, e["due"]); x.number_format = dfmt; x.fill = fill
        x = recs.cell(i, 6, e["late"]); x.fill = fill
        x = recs.cell(i, 7, f"=MAX({start_ref},E{i}+F{i})"); x.number_format = dfmt
        x = recs.cell(i, 8, f'=ROUND(IF(C{i}="payment plan",D{i}*(1-{plan_ref}),IF(C{i}="certain",D{i},D{i}*(1-{cut_ref}))),2)')
        x.number_format = cfmt
        x = recs.cell(i, 9, f'=IF(OR(C{i}="payment plan",C{i}="certain"),G{i},G{i}+{delay_ref})'); x.number_format = dfmt
        recs.cell(i, 10, e["source"])
    n_rec = max(2, len(res["receipts"]) + 1)
    for j, w in enumerate([34, 26, 15, 12, 13, 10, 16, 14, 14, 70], 1):
        recs.column_dimensions[COL(j)].width = w

    # Known Dates
    hdr = ["Line", "Category", "Amount", "Date money leaves", "Deadline", "Basis"]
    for j, h in enumerate(hdr, 1):
        known.cell(1, j, h).font = bold
    i = 1
    for e in res["payments"]:
        i += 1
        known.cell(i, 1, e["line"])
        known.cell(i, 2, e["category"])
        x = known.cell(i, 3, e["amount"]); x.number_format = cfmt; x.fill = fill
        x = known.cell(i, 4, e["date"]); x.number_format = dfmt; x.fill = fill
        if e.get("deadline"):
            known.cell(i, 5, e["deadline"]).number_format = dfmt
        known.cell(i, 6, ("[NEED] " if e.get("need") else "") + e["source"])
    if c:
        for e in res["commit_events"]:
            i += 1
            known.cell(i, 1, e["line"])
            known.cell(i, 2, "commitment")
            if e["k"] is None:
                x = known.cell(i, 3, f"={A['Commitment: one-off cost on the start date']}")
                known.cell(i, 4, f"={A['Commitment: start date']}").number_format = dfmt
            else:
                x = known.cell(i, 3, f"={A['Commitment: monthly cost']}")
                known.cell(i, 4, f"=EDATE({A['Commitment: first monthly payment']},{e['k']})").number_format = dfmt
            x.number_format = cfmt
            known.cell(i, 6, e["source"])
    n_known = max(2, i)
    for j, w in enumerate([34, 16, 12, 16, 13, 80], 1):
        known.column_dimensions[COL(j)].width = w

    rin, rout = list(res["base"]["lines_in"]), list(res["base"]["lines_out"])
    cells = {}

    def weekly(ws, case):
        title = "Normal quarter" if case == "base" else "Bad quarter (stress case)"
        ws["A1"] = f"{title}: week by week. Change the yellow cells on the schedules and Assumptions, not here."
        ws["A1"].font = bold
        ws.cell(2, 1, "Week").font = bold
        ws.cell(3, 1, "Week starts")
        ws.cell(4, 1, "Week ends")
        last = COL(1 + HORIZON_WEEKS)
        ws.cell(2, 2 + HORIZON_WEEKS, "13 weeks").font = bold
        for k in range(HORIZON_WEEKS):
            col = COL(2 + k)
            ws.cell(2, 2 + k, k + 1).font = bold
            ws.cell(3, 2 + k, f"={start_ref}+7*({col}2-1)").number_format = wfmt
            ws.cell(4, 2 + k, f"={col}3+6").number_format = wfmt
        r0 = 6
        ws.cell(r0, 1, "Opening balance").font = bold
        ws.cell(r0 + 1, 1, "Money in").font = bold
        r = r0 + 2
        amt_col, date_col = ("D", "G") if case == "base" else ("H", "I")
        first_in = r
        for ln in rin:
            ws.cell(r, 1, ln)
            for k in range(HORIZON_WEEKS):
                col = COL(2 + k)
                ws.cell(r, 2 + k, f"=SUMIFS('Receipts Schedule'!${amt_col}$2:${amt_col}${n_rec},'Receipts Schedule'!$A$2:$A${n_rec},$A{r},"
                                  f"'Receipts Schedule'!${date_col}$2:${date_col}${n_rec},\">=\"&{col}$3,"
                                  f"'Receipts Schedule'!${date_col}$2:${date_col}${n_rec},\"<=\"&{col}$4)").number_format = cfmt
            ws.cell(r, 2 + HORIZON_WEEKS, f"=SUM(B{r}:{last}{r})").number_format = cfmt
            r += 1
        tin = r
        ws.cell(tin, 1, "Total money in").font = bold
        r += 1
        ws.cell(r, 1, "Money out").font = bold
        r += 1
        first_out = r
        for ln in rout:
            ws.cell(r, 1, ln)
            for k in range(HORIZON_WEEKS):
                col = COL(2 + k)
                ws.cell(r, 2 + k, f"=SUMIFS('Known Dates'!$C$2:$C${n_known},'Known Dates'!$A$2:$A${n_known},$A{r},"
                                  f"'Known Dates'!$D$2:$D${n_known},\">=\"&{col}$3,'Known Dates'!$D$2:$D${n_known},\"<=\"&{col}$4)").number_format = cfmt
            ws.cell(r, 2 + HORIZON_WEEKS, f"=SUM(B{r}:{last}{r})").number_format = cfmt
            r += 1
        tout = r
        ws.cell(tout, 1, "Total money out").font = bold
        close, floor_r, head_r, below = tout + 1, tout + 2, tout + 3, tout + 4
        ws.cell(close, 1, "Closing balance").font = bold
        ws.cell(floor_r, 1, "Safety line")
        ws.cell(head_r, 1, "Above the safety line by")
        ws.cell(below, 1, "Below the line? (1 = yes)")
        for k in range(HORIZON_WEEKS):
            col, prev = COL(2 + k), COL(1 + k)
            ws.cell(r0, 2 + k, f"={A['Opening balance']}" if k == 0 else f"={prev}{close}").number_format = cfmt
            ws.cell(tin, 2 + k, f"=SUM({col}{first_in}:{col}{tin - 1})" if rin else 0).number_format = cfmt
            ws.cell(tout, 2 + k, f"=SUM({col}{first_out}:{col}{tout - 1})" if rout else 0).number_format = cfmt
            x = ws.cell(close, 2 + k, f"={col}{r0}+{col}{tin}-{col}{tout}"); x.number_format = cfmt; x.font = bold
            ws.cell(floor_r, 2 + k, f"={A['Safety line (the floor)']}").number_format = cfmt
            ws.cell(head_r, 2 + k, f"={col}{close}-{col}{floor_r}").number_format = cfmt
            ws.cell(below, 2 + k, f"=IF({col}{close}<{col}{floor_r},1,0)")
        ws.cell(tin, 2 + HORIZON_WEEKS, f"=SUM(B{tin}:{last}{tin})").number_format = cfmt
        ws.cell(tout, 2 + HORIZON_WEEKS, f"=SUM(B{tout}:{last}{tout})").number_format = cfmt
        crow = cbelow = None
        if c:
            crow, cclose, cbelow = below + 2, below + 3, below + 4
            ws.cell(crow, 1, f"Commitment: {c['label']}").font = bold
            ws.cell(cclose, 1, "Closing balance with the commitment").font = bold
            ws.cell(cbelow, 1, "Below the line with it? (1 = yes)")
            for k in range(HORIZON_WEEKS):
                col = COL(2 + k)
                ws.cell(crow, 2 + k, f"=SUMIFS('Known Dates'!$C$2:$C${n_known},'Known Dates'!$A$2:$A${n_known},$A{crow},"
                                     f"'Known Dates'!$D$2:$D${n_known},\">=\"&{col}$3,'Known Dates'!$D$2:$D${n_known},\"<=\"&{col}$4)").number_format = cfmt
                ws.cell(cclose, 2 + k, f"={col}{close}-SUM($B{crow}:{col}{crow})").number_format = cfmt
                ws.cell(cbelow, 2 + k, f"=IF({col}{cclose}<{col}{floor_r},1,0)")
            cells[case + "_cclose"] = cclose
        ws.column_dimensions["A"].width = 40
        for k in range(HORIZON_WEEKS + 1):
            ws.column_dimensions[COL(2 + k)].width = 11 if k < HORIZON_WEEKS else 13
        ws.freeze_panes = "B5"
        cells[case] = {"close": close, "below": below, "cbelow": cbelow}
    weekly(ws_base, "base")
    weekly(ws_st, "stress")

    # Summary
    headline_text = headline(m, res)
    summ["A1"] = f"13-week cash forecast: {m['business']}"
    summ["A1"].font = head
    summ["A2"] = f"Made {long_date(m['today'])}. Weeks run from {long_date(m['start'])} to {long_date(m['end'])}. Planning, not advice."
    last = COL(1 + HORIZON_WEEKS)
    bc, sc = cells["base"], cells["stress"]
    figs = [
        ("Opening balance", f"={A['Opening balance']}", cfmt),
        ("Balance date", f"={A['Balance date']}", dfmt),
        ("Safety line", f"={A['Safety line (the floor)']}", cfmt),
        ("Lowest balance, normal quarter", f"=MIN(Weekly!B{bc['close']}:{last}{bc['close']})", cfmt),
        ("Week of the lowest balance, normal quarter",
         f"=INDEX(Weekly!$B$2:${last}$2,MATCH(B7,Weekly!B{bc['close']}:{last}{bc['close']},0))", "0"),
        ("Lowest balance, bad quarter", f"=MIN(Stress!B{sc['close']}:{last}{sc['close']})", cfmt),
        ("Week of the lowest balance, bad quarter",
         f"=INDEX(Stress!$B$2:${last}$2,MATCH(B9,Stress!B{sc['close']}:{last}{sc['close']},0))", "0"),
        ("First week below the line, normal quarter",
         f'=IFERROR(INDEX(Weekly!$B$2:${last}$2,MATCH(1,Weekly!B{bc["below"]}:{last}{bc["below"]},0)),"none in 13 weeks")', "0"),
        ("First week below the line, bad quarter",
         f'=IFERROR(INDEX(Stress!$B$2:${last}$2,MATCH(1,Stress!B{sc["below"]}:{last}{sc["below"]},0)),"none in 13 weeks")', "0"),
        ("Bad-quarter runway (weeks before the line is crossed; 13 = no breach)",
         f"=IFERROR(MATCH(1,Stress!B{sc['below']}:{last}{sc['below']},0)-1,13)", "0"),
        ("Against the 8-week minimum", f'=IF(B13>={A["Bad-quarter runway the method asks for (weeks)"]},"meets it","below it")', None),
        ("Balance at week 13, normal quarter", f"=Weekly!{last}{bc['close']}", cfmt),
        ("Balance at week 13, bad quarter", f"=Stress!{last}{sc['close']}", cfmt),
    ]
    if c:
        figs.append(("First week below the line with the commitment, bad quarter",
                     f'=IFERROR(INDEX(Stress!$B$2:${last}$2,MATCH(1,Stress!B{sc["cbelow"]}:{last}{sc["cbelow"]},0)),"none in 13 weeks")', "0"))
    for i, (lab, f, fmt) in enumerate(figs, start=4):
        summ.cell(i, 1, lab)
        x = summ.cell(i, 2, f)
        x.font = bold
        if fmt:
            x.number_format = fmt
    r = 4 + len(figs) + 1
    for ln in reading_lines:
        summ.cell(r, 1, ln)
        summ.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        if ln and (ln == headline_text or not ln.startswith(("- ", "1.", "2.", "3.")) and len(ln) < 50):
            summ.cell(r, 1).font = bold
        summ.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        summ.row_dimensions[r].height = 15 * max(1, len(ln) // 80 + 1)
        r += 1
    summ.column_dimensions["A"].width = 64
    summ.column_dimensions["B"].width = 20
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "portrait" if ws.title == "Summary" else "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER if m.get("paper") == "Letter" else ws.PAPERSIZE_A4
    wb.save(path)
    return cells


def summary_lines(md):
    out = []
    for ln in md.splitlines()[2:]:
        if ln.startswith("|") or ln.startswith("# ") or ln.startswith("Made ") or ln == "## The figures":
            continue
        ln = ln.replace("**", "")
        out.append(ln[3:] if ln.startswith("## ") else ln)
    return out


def recalc_with_libreoffice(path):
    """Store calculated values beside the formulas so previews show figures. Optional: needs LibreOffice."""
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return False
    tmp = Path(tempfile.mkdtemp())
    try:
        subprocess.run([exe, f"-env:UserInstallation=file://{tmp / 'profile'}", "--headless", "--calc",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "out"), str(path)],
                       check=True, capture_output=True, timeout=180)
        made = tmp / "out" / Path(path).name
        if made.exists() and made.stat().st_size > 0:
            import shutil as sh
            sh.copyfile(made, path)
            return True
    except Exception:  # noqa: BLE001
        return False
    return False


def write_csv(folder, stem, m, res):
    import csv
    out = []
    for case, label in (("base", "normal"), ("stress", "bad-quarter")):
        p = folder / f"{stem}-{label}-weekly.csv"
        with p.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["Line"] + [f"Week {k['week']} ({k['start']})" for k in res[case]["weeks"]])
            w.writerow(["Opening balance"] + [k["opening"] for k in res[case]["weeks"]])
            for ln, v in res[case]["lines_in"].items():
                w.writerow([f"In: {ln}"] + v)
            w.writerow(["Total money in"] + [k["in"] for k in res[case]["weeks"]])
            for ln, v in res[case]["lines_out"].items():
                w.writerow([f"Out: {ln}"] + v)
            w.writerow(["Total money out"] + [k["out"] for k in res[case]["weeks"]])
            w.writerow(["Closing balance (= opening + money in - money out)"] + [k["closing"] for k in res[case]["weeks"]])
            w.writerow(["Safety line"] + [m["floor"]] * HORIZON_WEEKS)
        out.append(p)
    return out


# ------------------------------------------------------------------ the one-page reading (.docx, standard library)
def write_docx(path, md, title, paper="A4"):
    def runs(text):
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        xml = []
        for p_ in parts:
            if not p_:
                continue
            b = p_.startswith("**") and p_.endswith("**")
            t = p_[2:-2] if b else p_
            xml.append(f'<w:r><w:rPr>{"<w:b/>" if b else ""}<w:sz w:val="19"/></w:rPr>'
                       f'<w:t xml:space="preserve">{escape(t)}</w:t></w:r>')
        return "".join(xml)

    def para(text, size=None, bold=False, space_after=80):
        rpr = f'<w:rPr>{"<w:b/>" if bold else ""}<w:sz w:val="{size}"/></w:rPr>' if size else ""
        body = (f'<w:r>{rpr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>' if size else runs(text))
        return f'<w:p><w:pPr><w:spacing w:after="{space_after}"/></w:pPr>{body}</w:p>'
    body, table = [], []

    def flush():
        if not table:
            return
        rows_xml = []
        for ri, cells in enumerate(table):
            tcs = "".join(f'<w:tc><w:tcPr><w:tcW w:w="{3200 if ci == 0 else 3000}" w:type="dxa"/></w:tcPr>'
                          f'<w:p><w:pPr><w:spacing w:after="20"/></w:pPr>{runs(("**" + c + "**") if ri == 0 and c else c)}</w:p></w:tc>'
                          for ci, c in enumerate(cells))
            rows_xml.append(f"<w:tr>{tcs}</w:tr>")
        body.append('<w:tbl><w:tblPr><w:tblW w:w="9200" w:type="dxa"/><w:tblCellMar><w:left w:w="60" w:type="dxa"/>'
                    '<w:right w:w="60" w:type="dxa"/></w:tblCellMar></w:tblPr>' + "".join(rows_xml) + "</w:tbl>")
        body.append(para(""))
        table.clear()
    for ln in md.splitlines():
        if ln.startswith("|"):
            if re.match(r"^\|[-| ]+\|$", ln):
                continue
            table.append([c.strip() for c in ln.strip("|").split("|")])
            continue
        flush()
        if ln.startswith("# "):
            body.append(para(ln[2:], size=30, bold=True, space_after=60))
        elif ln.startswith("## "):
            body.append(para(ln[3:], size=22, bold=True, space_after=40))
        elif ln.strip():
            body.append(para(ln.replace("- ", "\u2022 ", 1) if ln.startswith("- ") else ln))
    flush()
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           + "".join(body) +
           '<w:sectPr>' + ('<w:pgSz w:w="12240" w:h="15840"/>' if paper == "Letter" else '<w:pgSz w:w="11906" w:h="16838"/>') + '<w:pgMar w:top="850" w:right="900" w:bottom="850" w:left="900" '
           'w:header="0" w:footer="0" w:gutter="0"/></w:sectPr></w:body></w:document>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
          '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
          '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
            '</Relationships>')
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>' + escape(title) + '</dc:title>'
            '<dc:creator>Made with Merigrove ' + SKILL_NAME + '</dc:creator></cp:coreProperties>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("docProps/core.xml", core)


# ------------------------------------------------------------------ CLI
def load(path):
    p = Path(path)
    if not p.exists():
        raise InputError(f"there is no file at {path}")
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise InputError(f"{path} is not valid JSON ({e.msg} at line {e.lineno})")


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40] or "business"


def print_compute(m, res, needs, notes):
    cur = m["currency"]
    bs, ss = res["base_stats"], res["stress_stats"]
    print(f"{m['business']} | made {m['today']} | weeks {m['start']} to {m['end']} | {cur}")
    print(f"OPENING: {money(m['opening'], cur, pence=True)} on {m['opening_date']} ({m['opening_source']})")
    print(f"FLOOR: {money(m['floor'], cur)} ({m['floor_basis']}); monthly outgoings {money(m['monthly_out'], cur)} = "
          + "; ".join(f"{k} {money(v, cur)}" for k, v in m["monthly_parts"]))
    print(f"STRESS: receipts -{round(m['cut'] * 100)}% and +{m['delay']} days; plan default {round(m['plan_default'] * 100)}%")
    print(f"LOCALE: country {m.get('country') or 'not set'}; paper {m.get('paper')}; dates {m.get('date_order') or 'not set'}; "
          f"decimals typed with {'a comma' if m.get('decimal_style') == 'comma' else 'a dot' if m.get('decimal_style') == 'dot' else 'no sign yet'}; "
          f"figures written like {money(1234567.89, cur, pence=True)}")
    pl = pattern_line(m)
    if pl:
        print("PATTERN: " + pl)
    for h in horizon_lines(m):
        print("HORIZON: " + h)
    print("WEEK | start | normal: in, out, closing | bad quarter: in, out, closing")
    for b, s in zip(res["base"]["weeks"], res["stress"]["weeks"]):
        print(f"  {b['week']:>2} | {b['start']} | {money(b['in'], cur)}, {money(b['out'], cur)}, {money(b['closing'], cur)}"
              f"{' BELOW' if b['below'] else ''} | {money(s['in'], cur)}, {money(s['out'], cur)}, {money(s['closing'], cur)}"
              f"{' BELOW' if s['below'] else ''}")
    print(f"NORMAL: lowest {money(bs['lowest'], cur)} in week {bs['lowest_week']}; first breach "
          f"{bs['first_breach'] or 'none'}; week 13 {money(bs['end'], cur)}")
    print(f"BAD QUARTER: lowest {money(ss['lowest'], cur)} in week {ss['lowest_week']}; first breach "
          f"{ss['first_breach'] or 'none'}; runway {ss['runway']} weeks (minimum {RUNWAY_MIN_WEEKS}); week 13 {money(ss['end'], cur)}")
    print("HEADLINE: " + headline(m, res))
    if decides(m, res):
        print("DECIDES: " + decides(m, res))
    if overdrawn_line(res):
        print("OVERDRAWN: " + overdrawn_line(res))
    for a in res["actions"]:
        print(f"ACTION: by {a['date']} | {a['line']} | {a['text']}")
    cl = commitment_line(m, res)
    if cl:
        print("COMMITMENT: " + cl)
    for t in known_lines_text(m, res):
        print("KNOWN: " + t)
    for x in notes:
        print("NOTE:", x)
    for x in needs:
        print("NEED:", x)


def results_payload(m, res, needs, notes, variance=None, suggestion=None):
    def ser(o):
        if isinstance(o, date):
            return o.isoformat()
        raise TypeError
    mm = {k: v for k, v in m.items() if k not in ("late_for", "history_daily", "tax_seen")}
    proj = [{"week_end": w["end"], "receipts": w["in"], "closing": w["closing"]} for w in res["base"]["weeks"]]
    return json.loads(json.dumps({"model": mm, "results": res, "needs": needs, "notes": notes, "projection": proj,
                                  "variance": variance, "suggestion": suggestion,
                                  "headline": headline(m, res), "decides": decides(m, res)}, default=ser))


def cmd_compute(a):
    m, needs, notes = normalise(load(a.inputs))
    res = compute(m)
    print_compute(m, res, needs, notes)
    sys.exit(1 if needs else 0)


def cmd_build(a):
    raw = load(a.inputs)
    if raw.get("provisional_tax") or raw.get("without_payer"):
        print("STOP: provisional_tax and without_payer are for the first read only. Put the founder's own figures in the "
              "inputs, remove both, and build again.")
        sys.exit(2)
    m, needs, notes = normalise(raw)
    res = compute(m)
    prev = load(a.previous) if a.previous else None
    var_txt, var = variance_text(m, prev)
    suggestion = stress_suggestion(prev, var)
    plain = reading_md(m, res, needs, notes, None, var_txt)
    note_text = None
    if a.note:
        if not Path(a.note).exists():
            raise InputError(f"there is no note file at {a.note}")
        note_text = Path(a.note).read_text(encoding="utf-8").strip()
        cur = m["currency"]
        figs = [money(e["amount"], cur) for e in res["receipts"] + res["payments"] + res["commit_events"]]
        figs += [money(e["s_amount"], cur) for e in res["receipts"]]
        figs += [money(w[k], cur) for case in ("base", "stress") for w in res[case]["weeks"] for k in ("in", "out", "closing")]
        allowed = plain + "\n" + " ".join(figs) + "\n" + "\n".join(notes)
        probs = check_note(note_text, allowed, m["currency"])
        if probs:
            for p in probs:
                print("PROBLEM:", p)
            print("STOP: the note was not used. Fix it and run build again. Nothing was written.")
            sys.exit(2)
    md = reading_md(m, res, needs, notes, note_text, var_txt)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    stem = f"cash-forecast-{slug(m['business'])}-{m['today'].isoformat()}"
    written = []
    try:
        fmt = a.format
        if fmt in ("auto", "xlsx"):
            try:
                import openpyxl  # noqa: F401
                p = out / f"{stem}.xlsx"
                written.append(p)
                write_xlsx(p, m, res, summary_lines(md))
                if not a.no_recalc and recalc_with_libreoffice(p):
                    print("NOTE: workbook recalculated, so previews show the figures as well as the formulas.")
                else:
                    print("NOTE: the workbook holds live formulas without stored values; Excel, Numbers, Sheets and "
                          "LibreOffice calculate them on opening, but a quick preview may show empty cells.")
            except ImportError:
                if fmt == "xlsx":
                    raise InputError("openpyxl is not installed here; run build with --format csv")
                fmt = "csv"
                print("NOTE: openpyxl is not available here, so the weekly tables are CSV files; the one-page reading carries every figure.")
        if fmt == "csv":
            written += write_csv(out, stem, m, res)
        pd = out / f"{stem}-one-page.docx"
        written.append(pd)
        write_docx(pd, md, f"13-week cash forecast: {m['business']}", m.get("paper") or "A4")
        rp = Path(a.results) if a.results else Path(a.inputs).with_name("results.json")
        rp.write_text(json.dumps(results_payload(m, res, needs, notes, var, suggestion), indent=2), encoding="utf-8")
        mdp = rp.with_name("reading.md")
        mdp.write_text(md, encoding="utf-8")
    except BaseException:
        for p in written:
            try:
                Path(p).unlink()
            except OSError:
                pass
        raise
    for p in written:
        print("WROTE:", p)
    print("RESULTS:", rp, "(scratch: for memory.py, not for the founder)")
    print("READING:", mdp)
    print("HEADLINE: " + headline(m, res))
    if decides(m, res):
        print("DECIDES: " + decides(m, res))
    if var_txt:
        print("SINCE LAST TIME: " + var_txt)
    if suggestion:
        print("SUGGEST: " + suggestion)
    for x in needs:
        print("NEED:", x)
    sys.exit(1 if needs else 0)


TEMPLATE = {
    "business": None, "currency": None, "date_order": None, "today": None,
    "opening": {"balance": None, "date": None, "source": None},
    "floor": {"mode": "one_month", "value": None},
    "known_outgoings": [{"label": None, "kind": "tax, payroll tax, government or other", "amount": None,
                         "date": None, "cadence": None, "next_date": None, "source": "who gave the figure"}],
    "recurring_out": [], "everyday_spending_week": {"amount": None}, "owner_pay_irregular_week": None,
    "recurring_in": [], "receipts_pattern_week": {"amount": None, "use": True},
    "invoices": [], "payment_terms_days": None, "instalments": [], "one_offs": [], "client_late_days": {},
    "commitment": None, "stress": {"receipts_cut": "20%", "delay_days": 15, "plan_default": "10%"},
    "_how": "Start from the draft find_recurring.py writes. Fill only what the founder gave; empty means a stated default or [NEED].",
}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd")
    sub.add_parser("template")
    c = sub.add_parser("compute")
    c.add_argument("inputs")
    b = sub.add_parser("build")
    b.add_argument("inputs")
    b.add_argument("--out", required=True)
    b.add_argument("--note")
    b.add_argument("--previous")
    b.add_argument("--results")
    b.add_argument("--format", default="auto", choices=["auto", "xlsx", "csv"])
    b.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args()
    if a.cmd == "template":
        print(json.dumps(TEMPLATE, indent=2))
    elif a.cmd == "compute":
        cmd_compute(a)
    elif a.cmd == "build":
        cmd_build(a)
    else:
        ap.print_help()
        sys.exit(2)


if __name__ == "__main__":
    try:
        main()
    except InputError as e:
        print(f"STOP: {e}. Nothing was written.")
        sys.exit(2)
    except AssertionError as e:
        print(f"STOP: the weekly arithmetic did not tie ({e}). Nothing usable was written; report this, do not show figures.")
        sys.exit(2)
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        print(f"STOP: the forecast could not finish ({type(e).__name__}: {e}). Nothing usable was written; "
              "check the inputs file and run again.")
        sys.exit(2)
